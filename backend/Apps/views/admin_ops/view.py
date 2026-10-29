#!/usr/bin/env python
# -*- coding: UTF-8 -*-
"""运营管理 API — 供独立 Web 运营台。

动作分组：
- login / me：运营登录与当前账号权限
- content：语录/主题/壁纸 CRUD、导入、枚举
- config：AppConfig 多语言×平台配置
- safety：安全词、危机热线、快捷开场、审计、内容举报
- membership：商品与 IAP 订单
- stats：DAU/留存/收入/地理等运营指标
- users：App 用户与运营账号
- rbac：角色权限矩阵
"""
import csv
import io
import json
import logging
from datetime import datetime, timedelta, timezone as dt_tz

from django.db import IntegrityError, transaction
from django.http import HttpResponse
from django.utils.translation import gettext as _
from drf_spectacular.utils import extend_schema
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.viewsets import ViewSet

from models.models import (
    Quote, Topic, Wallpaper, AppConfig, OnboardingQuestion,
    AiSafetyKeyword, CrisisHotline, AiQuickPrompt, User, IapReceipt, MembershipProduct,
    ContentReport, AiSafetyAudit, WelcomeScreen, Affirmation, AffirmationTemplate, AffirmationFollow,
    CollectionFolder, CollectionItem, InterstitialScreen, Quiz, QuizQuestion, AppIcon,
    DailyVerse, BibleTranslation, BibleBook, BibleChapter, BibleVerse,
    DevotionSeries, DevotionSession, MannaQuizLevel, MannaQuizQuestion, MannaQuizStory,
    PrayerContent, BadgeDef, MannaHomeCard, MannaTestimony, MannaStatsCopy,
    MannaContentTag, MannaLedger, UserLike, MannaFavorite, BibleHighlight, BibleNote,
    VerseCardTheme, WallpaperUsage, TagType,
)
from Apps.views.admin_ops.serializers import (
    QuoteAdminSerializer, WallpaperAdminSerializer, AdminLoginSerializer, ConfigSaveSerializer, ConfigPackDeleteSerializer,
    TopicSaveSerializer, WallpaperSaveSerializer, IdSerializer, QuoteSaveSerializer, BatchIdSerializer,
    as_tag_name, as_tag_names, normalize_wallpaper_media,
    OnboardingSaveSerializer, SafetyKeywordSaveSerializer, HotlineSaveSerializer,
    QuickPromptSaveSerializer, WelcomeScreenSaveSerializer, InterstitialScreenSaveSerializer,
    AppUserSaveSerializer,
    StaffUserSaveSerializer, RolePermissionsSaveSerializer, MembershipProductSaveSerializer,
    OrderStatusSerializer, ReportStatusSerializer, AffirmationTemplateSaveSerializer,
    QuizSaveSerializer, QuizQuestionSaveSerializer, AppIconSaveSerializer,
)
from tools.permissions import IsTokenValid, IsAdmin
from tools.selah_quota import (
    set_config, ensure_default_configs, ensure_default_welcome_screens,
    normalize_lang, normalize_platform, normalize_app_key,
)
from tools.tag_library import rebuild_wallpaper_index, rebuild_quote_index, ensure_tag_exists
from tools.token_tools import generate_login_user_token
from tools.auth_helpers import persist_auth
from tools.password_hasher import verify_password
from tools.utils import ApiResponse, CustomPagination
from tools.admin_rbac import (
    STAFF_ROLES, ROLE_LABELS, PERM_APP_USERS, PERM_STAFF, PERM_PERMISSIONS, PERM_OPS,
    PERM_CONTENT, has_perm, permissions_for_role, get_role_permissions_map, save_role_permissions,
    permission_catalog,
)
from tools.membership_products import (
    ensure_default_products, serialize_product, sync_iap_config_keys,
)
from tools.admin_scope import (
    apply_staff_app_list_filter, deny_staff_app_write, staff_managed_apps,
)
from tools.selah_enums import (
    content_enums_payload, validate_membership, INTEREST_TAGS, FAITH_TYPES, APP_LANGS,
    values_of, normalize_faith, interstitial_kind_values, all_onboarding_key_values,
    EXPLORE_TOPIC_CATEGORIES, TOPIC_CATEGORIES, TOPIC_TYPE_SLUGS,
)
from tools import admin_stats
from django.db.models import Q, Count, IntegerField, OuterRef, Subquery, Value
from django.db.models.functions import Coalesce
from django.utils import timezone as dj_tz
from django.utils.text import slugify


def _scope_from_query(request):
    """管理端精确维度：传了才过滤；空串=全局维。"""
    app = request.query_params.get('app_key')
    lang = request.query_params.get('language')
    plat = request.query_params.get('platform')
    return (
        None if app is None else normalize_app_key(app),
        None if lang is None else normalize_lang(lang),
        None if plat is None else normalize_platform(plat),
    )


def _parse_list_status(raw):
    """运营台列表 status：published/disabled 或 1/0；无法识别则不过滤。"""
    st = ('' if raw is None else str(raw)).strip().lower()
    if st in ('published', '1', 'true', 'yes'):
        return 1
    if st in ('disabled', '0', 'false', 'no'):
        return 0
    return None


def _json_tag_contains_q(field, name):
    """JSON 多值标签：纯字符串或 {name/value/label}（与 as_tag_names 输出对齐）。"""
    name = (name or '').strip()
    if not name:
        return Q()
    return (
        Q(**{f'{field}__contains': [name]})
        | Q(**{f'{field}__contains': [{'name': name}]})
        | Q(**{f'{field}__contains': [{'value': name}]})
        | Q(**{f'{field}__contains': [{'label': name}]})
    )


def _typed_tag_q(type_key, name, extra=None):
    """tag_index 的 type::name，可并上字段回落。"""
    name = (name or '').strip()
    if not type_key or not name:
        return Q()
    q = Q(tag_index__contains=[f'{type_key}::{name}'])
    if extra is not None:
        q |= extra
    return q


def _apply_scope_filters(qs, request, *, app_field='app_key'):
    """列表三维过滤 + Staff app_keys 落地（受限账号只看可管 App ∪ 空维）。"""
    app, lang, plat = _scope_from_query(request)
    allowed = staff_managed_apps(request.user)
    if app is not None:
        if allowed and app and app not in allowed:
            return qs.none()
        qs = qs.filter(**{app_field: app})
    elif allowed:
        qs = apply_staff_app_list_filter(qs, request.user, field=app_field)
    if lang is not None and hasattr(qs.model, 'language'):
        qs = qs.filter(language=lang)
    if plat is not None and hasattr(qs.model, 'platform'):
        qs = qs.filter(platform=plat)
    return qs


def _deny_app_write(request, app_key):
    """Staff 不可写目标 App 时返回 403 响应。"""
    msg = deny_staff_app_write(request.user, app_key)
    if msg:
        return ApiResponse(message=msg, code=403)
    return None


def _deny_empty_app_create(app_key):
    """内容新建禁止空维，避免双品牌回落串台。"""
    if not normalize_app_key(app_key or ''):
        return ApiResponse(message=_('请指定 App（禁止空维新建，避免双品牌串台）'), code=400)
    return None


def _refresh_devotion_session_count(series_id):
    """按启用课时重算系列 session_count。"""
    if not series_id:
        return
    DevotionSeries.objects.filter(id=series_id).update(
        session_count=DevotionSession.objects.filter(series_id=series_id, enabled=True).count()
    )


def _resolve_topic_slug(fields, existing=None):
    """类型主题用请求 slug；编辑未传则保留原值；新建再从 name 生成。"""
    slug = (fields.get('slug') or '').strip()
    if slug:
        return slug[:64]
    if existing is not None:
        old = (getattr(existing, 'slug', None) or '').strip()
        if old:
            return old[:64]
    name = (fields.get('name') or getattr(existing, 'name', '') or '')
    return (slugify(name) or 'topic')[:64]


_TOPIC_CATEGORY_KEYS = (
    set(values_of(EXPLORE_TOPIC_CATEGORIES))
    | set(values_of(TOPIC_CATEGORIES))
    | set(values_of(TOPIC_TYPE_SLUGS))
)


def _fill_topic_category(fields, existing=None):
    """运营台分类列读 category。类型主题前端常把 category 置空、只传 slug，这里回填。"""
    cat = (fields.get('category') or '').strip()
    if cat:
        fields['category'] = cat[:64]
        return
    slug = (fields.get('slug') or '').strip()
    if slug in _TOPIC_CATEGORY_KEYS:
        fields['category'] = slug
        return
    if existing is not None:
        fields['category'] = (getattr(existing, 'category', None) or '').strip()[:64]


_WALLPAPER_WRITE_FIELDS = (
    'image_url', 'media_type', 'video_url', 'poster_url', 'duration_sec',
    'is_vip', 'tags', 'name_en', 'name_es', 'name_pt', 'theme',
    'visual_tags', 'scenario_tags', 'emotion_tags', 'taxonomy', 'display_tags',
    'app_key', 'language', 'platform', 'sort_weight', 'status',
)
_WALLPAPER_TAG_FIELDS = ('theme', 'visual_tags', 'scenario_tags', 'emotion_tags', 'taxonomy')
_WALLPAPER_LIST_TAG_FIELDS = ('visual_tags', 'scenario_tags', 'emotion_tags', 'display_tags', 'tags')


def _normalize_quiz_bible_ref(raw):
    """题目经文锚点：book/chapter 必填；verse_start/verse_end 选填。

    返回 (dict, None) 或 (None, error_message)。空值统一为 {}。
    """
    if raw is None or raw == '' or raw == {}:
        return {}, None
    if not isinstance(raw, dict):
        return None, 'bible_ref must be object'
    book = str(raw.get('book') or '').strip().lower()
    chapter_raw = raw.get('chapter')
    vs_raw = raw.get('verse_start')
    ve_raw = raw.get('verse_end')
    has_any = bool(book) or chapter_raw not in (None, '') or vs_raw not in (None, '') or ve_raw not in (None, '')
    if not has_any:
        return {}, None
    if not book:
        return None, 'bible_ref.book required'
    try:
        chapter = int(chapter_raw)
    except (TypeError, ValueError):
        return None, 'bible_ref.chapter required'
    if chapter < 1:
        return None, 'bible_ref.chapter required'
    out = {'book': book, 'chapter': chapter}
    if vs_raw not in (None, ''):
        try:
            verse_start = int(vs_raw)
        except (TypeError, ValueError):
            return None, 'bible_ref.verse_start invalid'
        if verse_start < 1:
            return None, 'bible_ref.verse_start invalid'
        out['verse_start'] = verse_start
    if ve_raw not in (None, ''):
        if 'verse_start' not in out:
            return None, 'bible_ref.verse_start required when verse_end set'
        try:
            verse_end = int(ve_raw)
        except (TypeError, ValueError):
            return None, 'bible_ref.verse_end invalid'
        if verse_end < 1:
            return None, 'bible_ref.verse_end invalid'
        if verse_end < out['verse_start']:
            return None, 'bible_ref.verse_end must be >= verse_start'
        out['verse_end'] = verse_end
    return out, None


def _serialize_manna_story(r):
    return {
        'id': r.id,
        'title': r.title,
        'cover_url': r.cover_url or '',
        'card_url': r.card_url or '',
        'unlock_cost': r.unlock_cost,
        'knowledge_cards': r.knowledge_cards or [],
        'linked_level_id': r.linked_level_id,
        'body': r.body or '',
        'blocks': r.blocks or [],
        'cover_image_index': int(getattr(r, 'cover_image_index', 0) or 0),
        'sort_order': r.sort_order,
        'enabled': r.enabled,
        'app_key': r.app_key or '',
        'language': r.language or '',
        'platform': r.platform or '',
    }


def _apply_story_blocks_fields(patch, data, *, existing=None, require_image=False):
    """若请求带 blocks / cover_image_index，则校验并回写 cover_url/card_url/body。

    返回 error 字符串或 None。
    """
    from tools.manna_biz import normalize_story_blocks

    touched_blocks = 'blocks' in data
    touched_idx = 'cover_image_index' in data
    if not touched_blocks and not touched_idx:
        return None
    raw_blocks = data.get('blocks') if touched_blocks else (getattr(existing, 'blocks', None) if existing else [])
    raw_idx = data.get('cover_image_index') if touched_idx else (
        getattr(existing, 'cover_image_index', 0) if existing else 0
    )
    # 新建写 blocks 时要求至少 1 张图；仅改 index 或旧数据补 blocks 时按 require_image
    need_img = require_image if touched_blocks else False
    blocks, idx, cover_url, body, err = normalize_story_blocks(
        raw_blocks, raw_idx, require_image=need_img,
    )
    if err:
        return err
    patch['blocks'] = blocks
    patch['cover_image_index'] = idx
    patch['body'] = body
    if cover_url:
        patch['cover_url'] = cover_url
        patch['card_url'] = cover_url
    return None


def _payload_dict(request):
    """从 request.data 收集壁纸字段。不读 request.body，避免 DRF 已解析后 RawPostDataException。"""
    out = {}
    raw = getattr(request, 'data', None)
    if raw is None:
        return out
    try:
        for key in raw:
            dest = key[:-2] if isinstance(key, str) and key.endswith('[]') else key
            getter = getattr(raw, 'get', None)
            out[dest] = getter(key) if callable(getter) else raw[key]
    except TypeError:
        try:
            out.update(dict(raw))
        except (TypeError, ValueError):
            return out
    getlist = getattr(raw, 'getlist', None)
    if callable(getlist):
        try:
            for key in list(out.keys()):
                src = key if key in raw else f'{key}[]'
                try:
                    items = getlist(src)
                except Exception:
                    continue
                if src.endswith('[]') or len(items) > 1:
                    dest = src[:-2] if src.endswith('[]') else src
                    out[dest] = items
        except TypeError:
            pass
    return out


def _validation_error_text(exc):
    detail = getattr(exc, 'detail', None)
    if isinstance(detail, dict):
        parts = []
        for key, val in detail.items():
            if isinstance(val, (list, tuple)):
                val = val[0] if val else ''
            parts.append(f'{key}: {val}')
        return '; '.join(parts) or 'invalid'
    return str(detail or exc)


def _apply_wallpaper_media(fields, raw, validated, existing=None):
    """媒体字段有提交时，和已有行合并后写回 image/video/poster。只改 status 时不动。"""
    touched = any(key in raw for key in ('image_url', 'media_type', 'video_url', 'poster_url'))
    if existing is not None and not touched:
        return
    merged = {}
    for key in ('image_url', 'media_type', 'video_url', 'poster_url'):
        if key in raw:
            merged[key] = fields.get(key, raw.get(key))
        elif existing is not None:
            merged[key] = getattr(existing, key, '') or ''
        else:
            merged[key] = validated.get(key) or ''
    if 'duration_sec' in raw or (existing is None and 'duration_sec' in validated):
        merged['duration_sec'] = fields.get('duration_sec', validated.get('duration_sec'))
    normalized = normalize_wallpaper_media(merged, raw)
    for key in ('image_url', 'media_type', 'video_url', 'poster_url'):
        fields[key] = normalized[key]
    if 'duration_sec' in normalized:
        fields['duration_sec'] = normalized['duration_sec']


def _wallpaper_write_fields(request, validated):
    """只收集请求体里出现的壁纸可写字段（validated 优先，缺则回落原始 JSON）。"""
    raw = _payload_dict(request)
    out = {}
    for key in _WALLPAPER_WRITE_FIELDS:
        if key not in raw:
            continue
        if key in validated:
            out[key] = validated[key]
        else:
            out[key] = raw.get(key)
    if 'theme' in raw:
        out['theme'] = as_tag_name(validated.get('theme', raw.get('theme')))
    for key in _WALLPAPER_LIST_TAG_FIELDS:
        if key in raw:
            out[key] = as_tag_names(validated.get(key, raw.get(key)))
    return out


def _guard_scoped_write(request, *, existing=None, target_app='', require_app_key=False):
    """删除/更新/新建前的 Staff 校验；新建可强制非空 app_key。"""
    if existing is not None:
        denied = _deny_app_write(request, getattr(existing, 'app_key', '') or '')
        if denied:
            return denied
    denied = _deny_app_write(request, target_app)
    if denied:
        return denied
    if require_app_key:
        denied = _deny_empty_app_create(target_app)
        if denied:
            return denied
    return None


def _user_ids_for_scope_query(request):
    """按 app_key / language(system_lang) / platform 筛用户 id；无维条件返回 None。"""
    app = normalize_app_key(request.query_params.get('app_key') or '')
    lang = normalize_lang(
        request.query_params.get('language')
        or request.query_params.get('system_lang')
        or ''
    )
    plat = normalize_platform(request.query_params.get('platform') or '')
    if not app and not lang and not plat:
        return None
    qs = User.objects.filter(role='user')
    if app:
        qs = qs.filter(app_key=app)
    if lang:
        qs = qs.filter(system_lang=lang)
    if plat:
        qs = qs.filter(platform=plat)
    return list(qs.values_list('id', flat=True))

logger = logging.getLogger(__name__)



def _serialize_app_user(u):
    """App 端用户运营列表/详情序列化。"""
    return {
        'id': u.id,
        'username': u.username or '',
        'email': u.email or '',
        'device_id': u.device_id or '',
        'login_type': u.login_type,
        'display_name': u.display_name or '',
        'status': u.status,
        'vip_level': u.vip_level,
        'vip_expire_at': u.vip_expire_at.isoformat() if u.vip_expire_at else None,
        'is_trial': u.is_trial,
        'free_quota': u.free_quota,
        'vip_daily_remaining': u.vip_daily_remaining,
        'onboarding_completed': u.onboarding_completed,
        'faith_type': u.faith_type or '',
        'register_country': u.register_country or '',
        'system_lang': u.system_lang or '',
        'platform': getattr(u, 'platform', '') or '',
        'app_key': getattr(u, 'app_key', '') or '',
        'timezone': u.timezone or 'UTC',
        'checkin_streak': u.checkin_streak,
        'streak_hearts': u.streak_hearts,
        'newcomer_like_completed': u.newcomer_like_completed,
        'marketing_opt_in': u.marketing_opt_in,
        'profile_answers': getattr(u, 'profile_answers', None) or {},
        'reminder_prefs': getattr(u, 'reminder_prefs', None) or {},
        'widget_prefs': getattr(u, 'widget_prefs', None) or {},
        'alarm_prefs': getattr(u, 'alarm_prefs', None) or {},
        'prayer_lock_prefs': getattr(u, 'prayer_lock_prefs', None) or {},
        'selected_icon_key': getattr(u, 'selected_icon_key', None) or '',
        'checkin_record_enabled': u.checkin_record_enabled,
        'checkin_remind_enabled': u.checkin_remind_enabled,
        'created_at': u.created_at.isoformat() if u.created_at else None,
        'updated_at': u.updated_at.isoformat() if u.updated_at else None,
    }


def _serialize_staff(u):
    """运营后台账号序列化（含角色中文标签与可管理 App）。"""
    keys = getattr(u, 'app_keys', None) or []
    if not isinstance(keys, list):
        keys = []
    return {
        'id': u.id,
        'username': u.username or '',
        'email': u.email or '',
        'role': u.role,
        'role_label': ROLE_LABELS.get(u.role, u.role),
        'status': u.status,
        'app_keys': [str(k) for k in keys if k],
        'created_at': u.created_at.isoformat() if u.created_at else None,
        'updated_at': u.updated_at.isoformat() if u.updated_at else None,
    }


from Apps.views.admin_ops.mixins.recommend import RecommendMixin
from Apps.views.admin_ops.mixins.tag_library import TagLibraryMixin
from Apps.views.admin_ops.mixins.display_structure import DisplayStructureMixin


@extend_schema(tags=[_('运营后台')])
class AdminOpsViewSet(RecommendMixin, TagLibraryMixin, DisplayStructureMixin, ViewSet):
    """运营台聚合接口：内容/配置/安全/会员/统计/用户与 RBAC。"""

    def get_permissions(self):
        # login 免鉴权；其余必须登录 + 运营角色（IsAdmin）
        if self.action in ('login',):
            return []
        return [IsTokenValid(), IsAdmin()]

    @extend_schema(summary=_('运营登录'))
    @action(detail=False, methods=['post'], url_path='login')
    def login(self, request):
        """运营账号登录，返回 token 与角色权限。

        业务说明:
            仅 STAFF_ROLES 可登录运营台；App 端 role=user 拒绝。
            status=0 禁用；密码为空或校验失败统一「密码错误」（不区分是否有密码哈希）。
        """
        ser = AdminLoginSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        username = ser.validated_data['username']
        password = ser.validated_data['password']
        try:
            user = User.objects.get(username=username)
        except User.DoesNotExist:
            return ApiResponse(message=_('账号不存在'), code=400)
        # App 用户禁止进运营台
        if user.role not in STAFF_ROLES:
            return ApiResponse(message=_('非管理员'), code=403)
        if user.status == 0:
            return ApiResponse(message=_('账号已禁用'), code=403)
        if not user.password or not verify_password(password, user.password):
            return ApiResponse(message=_('密码错误'), code=400)
        token = generate_login_user_token(request, user, remember=True)
        perms = permissions_for_role(user.role)
        resp = ApiResponse(data={
            'token': token,
            'user_id': user.id,
            'username': user.username,
            'role': user.role,
            'permissions': perms,
            'app_keys': list(getattr(user, 'app_keys', None) or []),
        })
        persist_auth(request, resp, token, user.id, remember=True)
        return resp

    @extend_schema(summary=_('当前运营账号信息'))
    @action(detail=False, methods=['get'], url_path='me')
    def me(self, request):
        """当前运营账号与权限列表。"""
        user = request.user
        return ApiResponse(data={
            'user_id': user.id,
            'username': user.username,
            'role': user.role,
            'permissions': permissions_for_role(user.role),
            'app_keys': list(getattr(user, 'app_keys', None) or []),
        })

    @extend_schema(summary=_('内容枚举（运营台选项对齐）'))
    @action(detail=False, methods=['get'], url_path='content-enums')
    def content_enums(self, request):
        """内容枚举选项（运营台表单对齐）。

        业务说明:
            内容权限或运营权限任一即可（表单下拉与看板都可能用到）。
        """
        if not (has_perm(request.user, PERM_CONTENT) or has_perm(request.user, PERM_OPS)):
            return ApiResponse(message=_('无权限'), code=403)
        return ApiResponse(data=content_enums_payload())

    @extend_schema(summary=_('配置列表'))
    @action(detail=False, methods=['get'], url_path='configs')
    def configs(self, request):
        """列出 AppConfig（浏览筛或精确维）。

        业务说明:
            与内容列表一致：query 未传该维则不过滤；传空串=精确全局行。
            keys=a,b,c 只返回指定键。不做 App get_config 回退。
            meta.ai 暴露本机是否配置了 API Key（不回传密钥本身）。
        """
        import os
        ensure_default_configs()
        qs = AppConfig.objects.all().order_by('key', 'app_key', 'language', 'platform', '-priority', 'id')
        qs = _apply_scope_filters(qs, request)
        keys_raw = (request.query_params.get('keys') or '').strip()
        if keys_raw:
            key_list = [k.strip() for k in keys_raw.split(',') if k.strip()]
            if key_list:
                qs = qs.filter(key__in=key_list)
        app, lang, plat = _scope_from_query(request)
        data = [
            {
                'key': r.key,
                'value': r.value,
                'remark': r.remark,
                'app_key': r.app_key or '',
                'language': r.language or '',
                'platform': r.platform or '',
                'priority': int(getattr(r, 'priority', 0) or 0),
            }
            for r in qs
        ]
        return ApiResponse(
            data=data,
            meta={
                'app_key': '' if app is None else app,
                'language': '' if lang is None else lang,
                'platform': '' if plat is None else plat,
                'ai': {
                    'has_key': bool(os.getenv('AI_API_KEY')),
                    'api_base': os.getenv(
                        'AI_API_BASE',
                        'https://generativelanguage.googleapis.com/v1beta/openai',
                    ),
                    'model': (os.getenv('AI_MODEL') or '').strip() or None,
                },
                'translate': {
                    'has_key': bool(os.getenv('GOOGLE_TRANSLATE_API_KEY')),
                    'env_audio_root': bool((os.getenv('BIBLE_AUDIO_ROOT') or '').strip()),
                },
            },
        )

    @extend_schema(summary=_('保存配置'))
    @action(detail=False, methods=['post'], url_path='config-save')
    def config_save(self, request):
        """保存单条配置（指定 app/语言/平台/优先级维度）。

        业务说明:
            set_config 按 (key, app_key, language, platform, priority) upsert；空维写全局行。
            同维多 priority 时 App 单读取数字更大者。
        """
        ser = ConfigSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        app = normalize_app_key(ser.validated_data.get('app_key') or '')
        denied = _deny_app_write(request, app)
        if denied:
            return denied
        lang = normalize_lang(ser.validated_data.get('language') or '')
        plat = normalize_platform(ser.validated_data.get('platform') or '')
        pri = int(ser.validated_data.get('priority') or 0)
        set_config(
            ser.validated_data['key'],
            ser.validated_data['value'],
            remark=ser.validated_data.get('remark') or '',
            language=lang,
            platform=plat,
            app_key=app,
            priority=pri,
        )
        return ApiResponse(
            message='ok',
            data={'app_key': app, 'language': lang, 'platform': plat, 'priority': pri},
        )

    @extend_schema(summary=_('删除配置维度包'))
    @action(detail=False, methods=['post'], url_path='config-pack-delete')
    def config_pack_delete(self, request):
        """删除某精确维度下指定 keys 的 AppConfig 行。

        传入 priority 时只删该优先级；未传则删该维下全部优先级。
        """
        ser = ConfigPackDeleteSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        app = normalize_app_key(ser.validated_data.get('app_key') or '')
        denied = _deny_app_write(request, app)
        if denied:
            return denied
        lang = normalize_lang(ser.validated_data.get('language') or '')
        plat = normalize_platform(ser.validated_data.get('platform') or '')
        keys = list(ser.validated_data['keys'])
        qs = AppConfig.objects.filter(
            key__in=keys, app_key=app, language=lang, platform=plat,
        )
        pri = ser.validated_data.get('priority', None)
        if pri is not None:
            qs = qs.filter(priority=int(pri))
        deleted, _ = qs.delete()
        return ApiResponse(
            message='deleted',
            data={
                'deleted': deleted,
                'app_key': app,
                'language': lang,
                'platform': plat,
                'priority': pri,
            },
        )

    @extend_schema(summary=_('语录列表'))
    @action(detail=False, methods=['get'], url_path='quotes')
    def quotes(self, request):
        """语录分页列表（支持冷启池/关键词/状态/标签/回收站/app×语言×设备过滤）。

        业务说明:
            默认排除软删；trashed=1 只看回收站。
            cold 三态：未传=不过滤，1/true=仅冷启池，0/false=排除冷启池。
            status 接受 published/disabled 或 1/0；未传或 all=不过滤。
            tag_type + tag_name 按 tag_index（type::name）命中，并回落内容/主题/场景/情绪字段。
            q 对正文模糊搜；app_key/language/platform 与主题/壁纸一致：
            查询参数传了才精确过滤（含空串=全局维），未传则不过滤该维。
        """
        trashed = (request.query_params.get('trashed') or '').strip().lower() in ('1', 'true', 'yes')
        qs = Quote.objects.filter(is_deleted=trashed).order_by('-id')
        cold = request.query_params.get('in_cold_start_pool') or request.query_params.get('in_ops_pool')
        if cold in ('1', 'true', 'True'):
            qs = qs.filter(in_cold_start_pool=True)
        elif cold in ('0', 'false', 'False'):
            qs = qs.filter(in_cold_start_pool=False)
        q = (request.query_params.get('q') or '').strip()
        if q:
            qs = qs.filter(content__icontains=q)
        st = _parse_list_status(request.query_params.get('status'))
        if st is not None:
            qs = qs.filter(status=st)
        tag_type = (request.query_params.get('tag_type') or '').strip()
        tag_name = (request.query_params.get('tag_name') or '').strip()
        if tag_type and tag_name:
            tag_q = Q(tag_index__contains=[f'{tag_type}::{tag_name}'])
            if tag_type == 'content':
                tag_q |= Q(content_type=tag_name)
            elif tag_type == 'theme':
                tag_q |= Q(theme=tag_name)
            elif tag_type == 'scenario':
                tag_q |= _json_tag_contains_q('scenario_tags', tag_name)
            elif tag_type == 'emotion':
                tag_q |= _json_tag_contains_q('emotion_tags', tag_name)
            qs = qs.filter(tag_q)
        qs = _apply_scope_filters(qs, request)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        ser = QuoteAdminSerializer(page, many=True)
        return paginator.get_paginated_response(ser.data)

    @extend_schema(summary=_('语录保存'))
    @action(detail=False, methods=['post'], url_path='quote-save')
    def quote_save(self, request):
        """新建或局部更新语录。

        业务说明:
            更新时只 patch 请求体里实际出现的键，避免 Serializer default 把未传字段覆盖成空。
            content 校验对更新宽松：可不传 content 只改 is_vip 等开关；但若 id 不存在仍 404。
        """
        ser = QuoteSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        qid = data.pop('id', None)
        if qid:
            # 只更新请求体里出现的字段，避免 partial 覆盖成默认空值
            allowed = set(request.data.keys()) - {'id'}
            patch = {k: v for k, v in data.items() if k in allowed}
            existing = Quote.objects.filter(id=qid).first()
            if existing is None:
                return ApiResponse(message='not found', code=404)
            denied = _deny_app_write(request, existing.app_key or '')
            if denied:
                return denied
            if 'app_key' in patch:
                denied = _deny_app_write(request, patch.get('app_key') or '')
                if denied:
                    return denied
            # 防御性校验：最终状态为启用（status=1）时，content_type 和 theme 必填
            final_status = patch.get('status', existing.status)
            final_content_type = patch.get('content_type', existing.content_type)
            final_theme = patch.get('theme', existing.theme)
            if final_status == 1:
                missing = []
                if not (final_content_type or '').strip():
                    missing.append('content_type (文字内容类型)')
                if not (final_theme or '').strip():
                    missing.append('theme (主题标签)')
                if missing:
                    return ApiResponse(
                        message=f'启用前请先补充必填标签: {", ".join(missing)}',
                        code=400,
                    )
            if 'content' not in patch and 'content' not in allowed:
                pass
            # 重建 tag_index：合并 existing 与 patch 中的标签字段，确保归类统计正确
            if any(k in patch for k in ('content_type', 'theme', 'scenario_tags', 'emotion_tags', 'taxonomy')):
                merged_tags = {
                    'content_type': patch.get('content_type', existing.content_type or ''),
                    'theme': patch.get('theme', existing.theme or ''),
                    'scenario_tags': patch.get('scenario_tags', existing.scenario_tags or []),
                    'emotion_tags': patch.get('emotion_tags', existing.emotion_tags or []),
                    'taxonomy': patch.get('taxonomy', existing.taxonomy or {}),
                }
                rebuild_quote_index(merged_tags)
                patch['tag_index'] = merged_tags['tag_index']
            if patch:
                Quote.objects.filter(id=qid).update(**patch)
            # 自动创建缺失的隐性标签
            _final = {k: patch.get(k, getattr(existing, k, '') or '') for k in ('content_type', 'theme')}
            ensure_tag_exists('content', _final['content_type'])
            ensure_tag_exists('theme', _final['theme'])
            for _name in (patch.get('scenario_tags', existing.scenario_tags) or []):
                ensure_tag_exists('scenario', _name)
            for _name in (patch.get('emotion_tags', existing.emotion_tags) or []):
                ensure_tag_exists('emotion', _name)
            for _type_key, _names in (patch.get('taxonomy', existing.taxonomy) or {}).items():
                for _name in (_names or []):
                    ensure_tag_exists(_type_key, _name)
            return ApiResponse(data={'id': qid})
        denied = _deny_app_write(request, data.get('app_key') or '')
        if denied:
            return denied
        denied = _deny_empty_app_create(data.get('app_key') or '')
        if denied:
            return denied
        # 新建时如果是启用状态（status=1，默认），必须有 content_type 和 theme
        if data.get('status', 1) == 1:
            missing = []
            if not (data.get('content_type') or '').strip():
                missing.append('content_type (文字内容类型)')
            if not (data.get('theme') or '').strip():
                missing.append('theme (主题标签)')
            if missing:
                return ApiResponse(
                    message=f'启用前请先补充必填标签: {", ".join(missing)}',
                    code=400,
                )
        # 新建时重建 tag_index，确保归类统计正确
        rebuild_quote_index(data)
        q = Quote.objects.create(**data)
        # 自动创建缺失的隐性标签
        ensure_tag_exists('content', data.get('content_type', ''))
        ensure_tag_exists('theme', data.get('theme', ''))
        for _name in (data.get('scenario_tags') or []):
            ensure_tag_exists('scenario', _name)
        for _name in (data.get('emotion_tags') or []):
            ensure_tag_exists('emotion', _name)
        for _type_key, _names in (data.get('taxonomy') or {}).items():
            for _name in (_names or []):
                ensure_tag_exists(_type_key, _name)
        return ApiResponse(data={'id': q.id}, code=201)

    @extend_schema(summary=_('语录删除（支持批量）'))
    @action(detail=False, methods=['post'], url_path='quote-delete')
    def quote_delete(self, request):
        """软删除语录。

        业务说明:
            置 is_deleted=True，不物理删除，便于举报审计与恢复。
            支持 {"id": 1} 单删或 {"ids": [1,2,3]} 批量删。
        """
        ser = BatchIdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        ids = ser.get_ids()
        rows = Quote.objects.filter(id__in=ids)
        if rows.exists():
            for row in rows:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
            Quote.objects.filter(id__in=[r.id for r in rows]).update(is_deleted=True)
        return ApiResponse(message='ok', data={'deleted': len(rows)})

    @extend_schema(summary=_('从回收站恢复语录'))
    @action(detail=False, methods=['post'], url_path='quote-restore')
    def quote_restore(self, request):
        """恢复软删除语录（is_deleted=False）。"""
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        row = Quote.objects.filter(id=ser.validated_data['id']).first()
        if not row:
            return ApiResponse(message='not found', code=404)
        denied = _deny_app_write(request, row.app_key or '')
        if denied:
            return denied
        Quote.objects.filter(id=row.id).update(is_deleted=False)
        return ApiResponse(message='ok', data={'id': row.id})

    @extend_schema(summary=_('导出语录 CSV'))
    @action(detail=False, methods=['get'], url_path='quote-export')
    def quote_export(self, request):
        headers = [
            'id', 'content', 'language', 'source_ref', 'app_key', 'platform',
            'faith_tags', 'interest_tags', 'content_type', 'theme',
            'scenario_tags', 'emotion_tags', 'display_tags', 'taxonomy',
            'content_es', 'content_pt',
            'is_vip', 'countries',
            'audio_url', 'in_cold_start_pool', 'sort_weight', 'topic_id',
            'status', 'is_deleted', 'created_at', 'updated_at',
        ]
        def row(q):
            return [
                q.id, q.content, q.language, q.source_ref or '',
                q.app_key or '', q.platform or '',
                ','.join(q.faith_tags), ','.join(q.interest_tags),
                q.content_type or '', q.theme or '',
                ','.join(q.scenario_tags), ','.join(q.emotion_tags),
                ','.join(q.display_tags),
                json.dumps(q.taxonomy, ensure_ascii=False) if q.taxonomy else '',
                q.content_es or '', q.content_pt or '',
                'true' if q.is_vip else 'false',
                ','.join(q.countries),
                q.audio_url or '',
                'true' if q.in_cold_start_pool else 'false',
                q.sort_weight, q.topic_id or '',
                q.status, 'true' if q.is_deleted else 'false',
                q.created_at.strftime('%Y-%m-%d %H:%M:%S') if q.created_at else '',
                q.updated_at.strftime('%Y-%m-%d %H:%M:%S') if q.updated_at else '',
            ]
        return self._export_csv(
            request, Quote.objects.all(), headers, row,
            f'quotes_{datetime.now().strftime("%Y%m%d")}.csv',
        )

    @extend_schema(summary=_('语录导入模板下载'))
    @action(detail=False, methods=['get'], url_path='quote-import-template')
    def quote_import_template(self, request):
        """下载语录 CSV 导入模板，带 26 列表头 + 示例行，帮助运营正确对应标签列。"""
        headers = [
            'id', 'content', 'language', 'source_ref', 'app_key', 'platform',
            'faith_tags', 'interest_tags', 'content_type', 'theme',
            'scenario_tags', 'emotion_tags', 'display_tags', 'taxonomy',
            'content_es', 'content_pt',
            'is_vip', 'countries',
            'audio_url', 'in_cold_start_pool', 'sort_weight', 'topic_id',
            'status', 'is_deleted', 'created_at', 'updated_at',
        ]
        sample_row = [
            '',  # id 留空表示新建
            'Peace I leave with you; my peace I give you.',  # content
            'en',  # language
            'John 14:27',  # source_ref
            'selah',  # app_key
            '',  # platform 空=全端
            'new',  # faith_tags: new/old/other/learning
            'prayer,gratitude',  # interest_tags: 老枚举，10 个固定值 → prayer/family/anxiety/gratitude/courage/work/peace/hope/relationships/healing
            'bible_verse',  # content_type 单选（必填）
            'trust',  # theme 单选（必填）：自定义主题名，可随便填英文词（不要与 interest_tags 混淆）
            'before_sleep,morning_start',  # scenario_tags 多选逗号分隔
            'calm,hopeful',  # emotion_tags 多选逗号分隔
            '',  # display_tags 前端展示标签（显性）
            '',  # taxonomy JSON 对象字符串，如 {"season":"Advent"}
            'La paz os dejo; mi paz os doy.',  # content_es 西语可空
            '',  # content_pt 葡语可空
            'false',  # is_vip
            '',  # countries 空=全球
            '',  # audio_url 可空
            'true',  # in_cold_start_pool
            '0',  # sort_weight
            '',  # topic_id 可空
            '1',  # status 1上架/0下架
            'false',  # is_deleted
            '',  # created_at 导入时忽略
            '',  # updated_at 导入时忽略
        ]
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(headers)
        writer.writerow(sample_row)
        csv_bytes = buf.getvalue().encode('utf-8-sig')
        resp = HttpResponse(
            csv_bytes,
            content_type='text/csv; charset=utf-8',
        )
        resp['Content-Disposition'] = f'attachment; filename="quote_import_template.csv"'
        return resp

    @extend_schema(summary=_('Excel/CSV 导入语录'))
    @action(detail=False, methods=['post'], url_path='quote-import')
    def quote_import(self, request):
        """支持 csv_text 文本，或 multipart file（.csv / .xlsx）。

        业务说明:
            - xlsx 用 openpyxl 读首 sheet；首行作 header（空格→下划线）。
            - 标签非法或 language 不在白名单 → 记入 errors 并跳过该行。
            - 全部失败且 imported=0 → 400；部分成功仍 200，errors 最多回 20 条。
            - countries 逗号分隔；空=全球可见。
        """
        rows = []
        upload = request.FILES.get('file')
        if upload:
            name = (getattr(upload, 'name', None) or '').lower()
            raw = upload.read()
            # PK zip 头：即使扩展名错也可识别 xlsx
            is_xlsx = (
                name.endswith('.xlsx')
                or name.endswith('.xlsm')
                or raw[:2] == b'PK'
            )
            if is_xlsx:
                try:
                    from openpyxl import load_workbook
                except ImportError:
                    return ApiResponse(message='openpyxl required for Excel import', code=500)
                wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
                ws = wb.active
                header = None
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    vals = [('' if v is None else str(v)).strip() for v in row]
                    if i == 0:
                        header = [h.lower().replace(' ', '_') for h in vals]
                        continue
                    if not any(vals):
                        continue
                    rows.append(dict(zip(header, vals)))
            else:
                text = raw.decode('utf-8-sig', errors='ignore')
                rows = list(csv.DictReader(io.StringIO(text)))
        else:
            text = request.data.get('csv_text') or ''
            if not text:
                return ApiResponse(message='csv_text or file required', code=400)
            rows = list(csv.DictReader(io.StringIO(text)))

        count = 0
        skipped_count = 0
        errors = []
        skipped = []
        allowed_langs = set(values_of(APP_LANGS))
        raw_default = ''
        if hasattr(request.data, 'get'):
            raw_default = request.data.get('app_key') or ''
        default_app = normalize_app_key(raw_default)
        # TagType.key → Quote 字段名 映射（scope=text/both 适用）
        _TAGTYPE_KEY_TO_QUOTE_FIELD = {
            'content': 'content_type',
            'theme': 'theme',
            'scenario': 'scenario_tags',
            'emotion': 'emotion_tags',
        }
        # 必填标签字段列表（scope=text/both 且 required=True，且能映射到 Quote 的）
        _required_fields = [
            _TAGTYPE_KEY_TO_QUOTE_FIELD[k]
            for k in TagType.objects.filter(required=True, scope__in=('text', 'both'))
                .values_list('key', flat=True).distinct()
            if k in _TAGTYPE_KEY_TO_QUOTE_FIELD
        ]
        # start=2：与表格行号对齐（第 1 行是表头）
        for idx, row in enumerate(rows, start=2):
            content = (row.get('content') or '').strip()
            if not content:
                continue
            tags = [t.strip() for t in (row.get('interest_tags') or '').split(',') if t.strip()]
            faith = [normalize_faith(t.strip()) for t in (row.get('faith_tags') or '').split(',') if t.strip()]
            ok_tags, bad_tags = validate_membership(tags, INTEREST_TAGS)
            ok_faith, bad_faith = validate_membership(faith, FAITH_TYPES, field_name='faith_tags')
            lang = (row.get('language') or 'en').strip() or 'en'
            plat = normalize_platform(row.get('platform') or '')
            app = normalize_app_key(row.get('app_key') or default_app)
            empty_msg = None
            if not app:
                empty_msg = _('请指定 App（禁止空维新建，避免双品牌串台）')
            staff_msg = deny_staff_app_write(request.user, app) if app else None
            if empty_msg or staff_msg:
                errors.append({
                    'row': idx,
                    'app_key': app or '(empty)',
                    'error': empty_msg or staff_msg,
                })
                continue
            # 标签或语言非法：跳过该行，继续导入其余行
            if bad_tags or bad_faith or lang not in allowed_langs:
                errors.append({
                    'row': idx,
                    'interest_tags': bad_tags,
                    'faith_tags': bad_faith,
                    'language': None if lang in allowed_langs else lang,
                })
                continue
            # ---------- 新字段解析 ----------
            content_type = (row.get('content_type') or '').strip()
            theme = (row.get('theme') or '').strip()
            content_es = (row.get('content_es') or '').strip()
            content_pt = (row.get('content_pt') or '').strip()
            scenario_tags = [t.strip() for t in (row.get('scenario_tags') or '').split(',') if t.strip()]
            emotion_tags = [t.strip() for t in (row.get('emotion_tags') or '').split(',') if t.strip()]
            display_tags = [t.strip() for t in (row.get('display_tags') or '').split(',') if t.strip()]
            # taxonomy 是 dict 类型，尝试 JSON 解析（空或非法则为 {}）
            taxonomy_raw = (row.get('taxonomy') or '').strip()
            taxonomy = {}
            if taxonomy_raw:
                try:
                    parsed = json.loads(taxonomy_raw)
                    taxonomy = parsed if isinstance(parsed, dict) else {}
                except (json.JSONDecodeError, TypeError):
                    taxonomy = {}
            # rebuild_quote_index 需要 value 是 list 才不会逐字符迭代 str
            taxonomy_for_index = {
                k: (v if isinstance(v, list) else [v])
                for k, v in taxonomy.items()
                if isinstance(v, (list, str))
            }
            # ---------- 必填标签校验（TagType.required=True 的 scope=text/both 字段）----------
            missing_required = []
            _parsed_values = {
                'content_type': content_type,
                'theme': theme,
                'scenario_tags': scenario_tags,
                'emotion_tags': emotion_tags,
            }
            for field_name in _required_fields:
                raw = _parsed_values.get(field_name)
                if field_name.endswith('_tags'):
                    # 多选字段：list 为空
                    if not raw:
                        missing_required.append(field_name)
                else:
                    # 单选字段：字符串为空
                    if not raw:
                        missing_required.append(field_name)
            if missing_required:
                errors.append({
                    'row': idx,
                    'missing_required': missing_required,
                    'error': '必填标签缺失（见 TagType.required）',
                })
                continue
            tid_raw = (row.get('topic_id') or '').strip()
            topic_id = int(tid_raw) if tid_raw.isdigit() else None
            raw_id = str(row.get('id') or '').strip()
            row_id = int(raw_id) if raw_id.isdigit() else None

            defaults = dict(
                content=content,
                language=lang,
                source_ref=(row.get('source_ref') or '').strip(),
                app_key=app,
                platform=plat,
                faith_tags=ok_faith,
                interest_tags=ok_tags,
                content_type=content_type,
                theme=theme,
                content_es=content_es,
                content_pt=content_pt,
                scenario_tags=scenario_tags,
                emotion_tags=emotion_tags,
                display_tags=display_tags,
                taxonomy=taxonomy,
                is_vip=str(row.get('is_vip') or '').lower() in ('1', 'true', 'yes'),
                countries=[c.strip() for c in (row.get('countries') or '').split(',') if c.strip()],
                audio_url=(row.get('audio_url') or '').strip(),
                in_cold_start_pool=str(
                    row.get('in_cold_start_pool') or row.get('in_ops_pool') or ''
                ).lower() in ('1', 'true', 'yes'),
                sort_weight=int(str(row.get('sort_weight') or 0).strip() or 0),
                topic_id=topic_id,
                status=int(str(row.get('status') or 1).strip() or 1),
                is_deleted=str(row.get('is_deleted') or '').lower() in ('1', 'true', 'yes'),
            )
            # 自动重建 tag_index（根据 content_type/theme/scenario/emotion/taxonomy 生成检索键）
            # 注意：rebuild_quote_index 只接受 value 可迭代的 taxonomy 项，完整 taxonomy 不丢失
            _saved_tax = defaults['taxonomy']
            defaults['taxonomy'] = taxonomy_for_index
            rebuild_quote_index(defaults)
            defaults['taxonomy'] = _saved_tax
            # Upsert：有 id 且存在则更新，否则新建
            # 自动创建缺失的隐性标签
            ensure_tag_exists('content', defaults.get('content_type', ''))
            ensure_tag_exists('theme', defaults.get('theme', ''))
            for _name in (defaults.get('scenario_tags') or []):
                ensure_tag_exists('scenario', _name)
            for _name in (defaults.get('emotion_tags') or []):
                ensure_tag_exists('emotion', _name)
            for _type_key, _names in (defaults.get('taxonomy') or {}).items():
                for _name in (_names if isinstance(_names, list) else [_names] if _names else []):
                    ensure_tag_exists(_type_key, _name)
            if row_id:
                existing = Quote.objects.filter(id=row_id).first()
                if existing:
                    for k, v in defaults.items():
                        setattr(existing, k, v)
                    existing.save()
                    count += 1
                    continue
            # 新建前先按 content + app_key + language 去重（避免重复上传产生重复记录）
            # 同 wallpaper_import 的 image_url 去重策略
            dup_qs = Quote.objects.filter(
                content=content, app_key=app, language=lang, is_deleted=False
            )
            if row_id:
                dup_qs = dup_qs.exclude(id=row_id)
            dup = dup_qs.first()
            if dup:
                errors.append({
                    'row': idx,
                    'error': f'内容重复，已存在于 id={dup.id}（同 content+app_key+language），跳过',
                })
                continue
            # 新建路径：强制 is_deleted=False（导入不应新建软删数据）
            defaults['is_deleted'] = False
            Quote.objects.create(**defaults)
            count += 1
        logger.info('quote_import count=%s errors=%s', count, len(errors))
        # 全部失败才 400；部分成功仍返回 imported + errors 样例
        if errors and not count:
            return ApiResponse(
                message='import failed: invalid tags/language/app_key',
                code=400,
                data={'imported': 0, 'errors': errors[:20]},
            )
        return ApiResponse(data={'imported': count, 'errors': errors[:20]})

    # ---------- Export helpers ----------
    def _export_csv(self, request, queryset, headers, row_func, filename):
        """通用 CSV 导出：queryset 按 row_func 转 list，headers 做表头。"""
        denied = self._require_content(request)
        if denied:
            return denied
        if hasattr(request, 'query_params'):
            app = normalize_app_key(request.query_params.get('app_key') or '')
            lang = request.query_params.get('language') or ''
            plat = normalize_platform(request.query_params.get('platform') or '')
        else:
            app = normalize_app_key(request.GET.get('app_key') or '')
            lang = request.GET.get('language') or ''
            plat = normalize_platform(request.GET.get('platform') or '')
        if app:
            queryset = queryset.filter(app_key=app)
        if lang:
            queryset = queryset.filter(language=lang)
        if plat:
            queryset = queryset.filter(platform=plat)
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(headers)
        for obj in queryset.order_by('id'):
            writer.writerow(row_func(obj))
        csv_bytes = buf.getvalue().encode('utf-8-sig')
        resp = HttpResponse(
            csv_bytes,
            content_type='text/csv; charset=utf-8',
        )
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        return resp

    @extend_schema(summary=_('导出壁纸 CSV'))
    @action(detail=False, methods=['get'], url_path='wallpaper-export')
    def wallpaper_export(self, request):
        headers = [
            'id', 'image_url', 'media_type', 'video_url', 'poster_url', 'duration_sec',
            'is_vip', 'name_en', 'name_es', 'name_pt', 'theme',
            'visual_tags', 'scenario_tags', 'emotion_tags',
            'app_key', 'language', 'platform', 'sort_weight',
            'status', 'is_deleted', 'created_at', 'updated_at',
        ]
        def row(w):
            return [
                w.id, w.image_url, w.media_type, w.video_url, w.poster_url, w.duration_sec or '',
                'true' if w.is_vip else 'false', w.name_en, w.name_es, w.name_pt, w.theme,
                ','.join(w.visual_tags), ','.join(w.scenario_tags), ','.join(w.emotion_tags),
                w.app_key or '', w.language or '', w.platform or '', w.sort_weight,
                w.status, 'true' if w.is_deleted else 'false',
                w.created_at.strftime('%Y-%m-%d %H:%M:%S') if w.created_at else '',
                w.updated_at.strftime('%Y-%m-%d %H:%M:%S') if w.updated_at else '',
            ]
        return self._export_csv(
            request, Wallpaper.objects.all(), headers, row,
            f'wallpapers_{datetime.now().strftime("%Y%m%d")}.csv',
        )

    @extend_schema(summary=_('Excel/CSV 导入壁纸'))
    @action(detail=False, methods=['post'], url_path='wallpaper-import')
    def wallpaper_import(self, request):
        """批量导入壁纸。支持 CSV / Excel 文件，或纯 JSON wallpapers 数组。

        CSV 表头：image_url（静态必填）, name_en, name_es, name_pt, theme,
                  visual_tags, scenario_tags, emotion_tags,
                  is_vip, media_type, video_url, poster_url, sort_weight,
                  app_key, language, platform
        视频行要有 video_url，且 poster_url 必须是图片。静态行会清空 video_url 和 poster_url。
        标签字段用英文逗号分隔，如 "Nature,Peace"。
        """
        denied = self._require_content(request)
        if denied:
            return denied
        rows = []
        upload = request.FILES.get('file')
        if upload:
            name = (getattr(upload, 'name', None) or '').lower()
            raw = upload.read()
            is_xlsx = name.endswith('.xlsx') or name.endswith('.xlsm') or raw[:2] == b'PK'
            if is_xlsx:
                try:
                    from openpyxl import load_workbook
                except ImportError:
                    return ApiResponse(message='openpyxl required for Excel import', code=500)
                wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
                ws = wb.active
                header = None
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    vals = [('' if v is None else str(v)).strip() for v in row]
                    if i == 0:
                        header = [h.lower().replace(' ', '_') for h in vals]
                        continue
                    if not any(vals):
                        continue
                    rows.append(dict(zip(header, vals)))
            else:
                text = raw.decode('utf-8-sig', errors='ignore')
                rows = list(csv.DictReader(io.StringIO(text)))
        else:
            data = request.data if isinstance(request.data, dict) else {}
            json_rows = data.get('wallpapers') or data.get('rows')
            if isinstance(json_rows, list) and json_rows:
                rows = json_rows
            else:
                text = data.get('csv_text') or ''
                if not text:
                    return ApiResponse(message='file / csv_text / wallpapers[] 至少提供一种', code=400)
                rows = list(csv.DictReader(io.StringIO(text)))

        count = 0
        skipped_count = 0
        errors = []
        skipped = []
        raw_default = request.data.get('app_key') if hasattr(request.data, 'get') else ''
        default_app = normalize_app_key(raw_default)

        def _parse_tag(raw):
            if isinstance(raw, list):
                return [str(t).strip() for t in raw if str(t).strip()]
            return [t.strip() for t in str(raw or '').split(',') if t.strip()]

        for idx, row in enumerate(rows, start=2):
            raw_id = str(row.get('id') or '').strip()
            row_id = int(raw_id) if raw_id.isdigit() else None
            name_en = (row.get('name_en') or '').strip()
            name_es = (row.get('name_es') or '').strip()
            name_pt = (row.get('name_pt') or '').strip()
            theme = (row.get('theme') or '').strip()
            try:
                media = normalize_wallpaper_media({
                    'image_url': (row.get('image_url') or '').strip(),
                    'media_type': (row.get('media_type') or '').strip(),
                    'video_url': (row.get('video_url') or '').strip(),
                    'poster_url': (row.get('poster_url') or '').strip(),
                    'duration_sec': (row.get('duration_sec') or '').strip() or None,
                }, row)
            except ValidationError as exc:
                errors.append({'row': idx, 'error': _validation_error_text(exc)})
                continue
            image_url = media['image_url']
            media_type = media['media_type']
            video_url = media['video_url']
            poster_url = media['poster_url']
            is_vip = str(row.get('is_vip') or '').lower() in ('1', 'true', 'yes')
            sort_raw = str(row.get('sort_weight') or '').strip()
            sort_weight = int(sort_raw) if sort_raw.lstrip('-').isdigit() else 0
            plat = normalize_platform(row.get('platform') or '')
            app = normalize_app_key(row.get('app_key') or default_app)
            lang = (row.get('language') or '').strip()

            denied = deny_staff_app_write(request.user, app) if app else None
            if denied:
                errors.append({'row': idx, 'app_key': app, 'error': str(denied)})
                continue

            defaults = dict(
                image_url=image_url, media_type=media_type,
                video_url=video_url, poster_url=poster_url,
                is_vip=is_vip, theme=theme,
                visual_tags=_parse_tag(row.get('visual_tags')),
                scenario_tags=_parse_tag(row.get('scenario_tags')),
                emotion_tags=_parse_tag(row.get('emotion_tags')),
                name_en=name_en, name_es=name_es, name_pt=name_pt,
                app_key=app, language=lang, platform=plat,
                sort_weight=sort_weight,
                status=1, is_deleted=False,
            )
            # 自动创建缺失的隐性标签
            ensure_tag_exists('theme', defaults.get('theme', ''))
            for _name in (defaults.get('visual_tags') or []):
                ensure_tag_exists('visual', _name)
            for _name in (defaults.get('scenario_tags') or []):
                ensure_tag_exists('scenario', _name)
            for _name in (defaults.get('emotion_tags') or []):
                ensure_tag_exists('emotion', _name)
            existing = None
            if row_id:
                existing = Wallpaper.objects.filter(id=row_id).first()
            if existing:
                for k, v in defaults.items():
                    setattr(existing, k, v)
                existing.save()
                count += 1
            else:
                dup = Wallpaper.objects.filter(image_url=image_url).first()
                if dup:
                    skipped_count += 1
                    skipped.append({'row': idx, 'image_url': image_url,
                                    'reason': f'already exists (id={dup.id})'})
                    continue
                Wallpaper.objects.create(**defaults)
                count += 1

        logger.info('wallpaper_import count=%s skipped=%s errors=%s', count, skipped_count, len(errors))
        if errors and not count:
            return ApiResponse(
                message='import failed: check image_url / app_key',
                code=400,
                data={'imported': 0, 'skipped': skipped[:20], 'errors': errors[:20]},
            )
        return ApiResponse(data={'imported': count, 'skipped': skipped[:20], 'errors': errors[:20]})

    @extend_schema(summary=_('Onboarding 题目'))
    @action(detail=False, methods=['get', 'post'], url_path='onboarding-questions')
    def onboarding_questions(self, request):
        """引导题列表或保存/删除。

        业务说明:
            GET：language/platform 查询参数「传了才过滤」（含空串=全局维度）。
            POST：delete_id 硬删；有 id 则按请求体键 patch；无 id 新建。
        """
        if request.method == 'GET':
            qs = OnboardingQuestion.objects.all().order_by('sort_order', 'id')
            qs = _apply_scope_filters(qs, request)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([
                {
                    'id': r.id, 'question_key': r.question_key, 'title': r.title,
                    'options': r.options, 'sort_order': r.sort_order, 'enabled': r.enabled,
                    'app_key': r.app_key or '',
                    'language': r.language or '', 'platform': r.platform or '',
                }
                for r in page
            ])
        ser = OnboardingSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        if data.get('delete_id'):
            row = OnboardingQuestion.objects.filter(id=data['delete_id']).first()
            if row:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
                row.delete()
            return ApiResponse(message='deleted')
        if 'app_key' in data:
            data['app_key'] = normalize_app_key(data.get('app_key') or '')
        if 'language' in data:
            data['language'] = normalize_lang(data.get('language') or '')
        if 'platform' in data:
            data['platform'] = normalize_platform(data.get('platform') or '')
        qid = data.pop('id', None)
        data.pop('delete_id', None)
        if qid:
            existing = OnboardingQuestion.objects.filter(id=qid).first()
            if not existing:
                return ApiResponse(message='not found', code=404)
            denied = _deny_app_write(request, existing.app_key or '')
            if denied:
                return denied
            patch = {k: v for k, v in data.items() if k in request.data}
            if 'app_key' in patch:
                denied = _deny_app_write(request, patch.get('app_key') or '')
                if denied:
                    return denied
            if patch:
                try:
                    OnboardingQuestion.objects.filter(id=qid).update(**patch)
                except IntegrityError:
                    return ApiResponse(message=_('同维度已存在相同 question_key'), code=400)
            return ApiResponse(data={'id': qid})
        denied = _guard_scoped_write(
            request, target_app=data.get('app_key') or '', require_app_key=True,
        )
        if denied:
            return denied
        try:
            r = OnboardingQuestion.objects.create(**data)
        except IntegrityError:
            return ApiResponse(message=_('同维度已存在相同 question_key'), code=400)
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('安全词'))
    @action(detail=False, methods=['get', 'post'], url_path='safety-keywords')
    def safety_keywords(self, request):
        """安全关键词列表或保存/删除。

        业务说明:
            与引导题相同：GET 按维度过滤；POST 支持 delete_id / id patch / 新建。
            category 须在 SAFETY_CATEGORIES 白名单（Serializer 校验）。
        """
        if request.method == 'GET':
            qs = AiSafetyKeyword.objects.all().order_by('category', 'id')
            qs = _apply_scope_filters(qs, request)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([
                {
                    'id': r.id, 'category': r.category, 'keyword': r.keyword, 'enabled': r.enabled,
                    'app_key': r.app_key or '',
                    'language': r.language or '', 'platform': r.platform or '',
                }
                for r in page
            ])
        ser = SafetyKeywordSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = ser.validated_data
        if data.get('delete_id'):
            row = AiSafetyKeyword.objects.filter(id=data['delete_id']).first()
            if row:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
                row.delete()
            return ApiResponse(message='deleted')
        app = normalize_app_key(data.get('app_key') or '')
        lang = normalize_lang(data.get('language') or '')
        plat = normalize_platform(data.get('platform') or '')
        if data.get('id'):
            existing = AiSafetyKeyword.objects.filter(id=data['id']).first()
            if existing:
                denied = _deny_app_write(request, existing.app_key or '')
                if denied:
                    return denied
            denied = _deny_app_write(request, app)
            if denied:
                return denied
            patch = {}
            for k in ('category', 'keyword', 'enabled', 'language', 'platform', 'app_key'):
                if k in request.data:
                    if k == 'language':
                        patch[k] = lang
                    elif k == 'platform':
                        patch[k] = plat
                    elif k == 'app_key':
                        patch[k] = app
                    else:
                        patch[k] = data.get(k)
            if patch:
                AiSafetyKeyword.objects.filter(id=data['id']).update(**patch)
            return ApiResponse(data={'id': data['id']})
        denied = _guard_scoped_write(request, target_app=app, require_app_key=True)
        if denied:
            return denied
        r = AiSafetyKeyword.objects.create(
            category=data['category'],
            keyword=data['keyword'],
            enabled=bool(data.get('enabled', True)),
            app_key=app,
            language=lang,
            platform=plat,
        )
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('危机热线'))
    @action(detail=False, methods=['get', 'post'], url_path='hotlines')
    def hotlines(self, request):
        """危机热线列表或保存/删除。

        业务说明:
            country 供 AI 自残拒答时按注册国匹配；GLOBAL 作兜底。
        """
        if request.method == 'GET':
            qs = CrisisHotline.objects.all().order_by('id')
            qs = _apply_scope_filters(qs, request)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([
                {
                    'id': r.id, 'country': r.country, 'label': r.label,
                    'phone_or_url': r.phone_or_url, 'enabled': r.enabled,
                    'app_key': r.app_key or '',
                    'language': r.language or '', 'platform': r.platform or '',
                }
                for r in page
            ])
        ser = HotlineSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = ser.validated_data
        if data.get('delete_id'):
            row = CrisisHotline.objects.filter(id=data['delete_id']).first()
            if row:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
                row.delete()
            return ApiResponse(message='deleted')
        app = normalize_app_key(data.get('app_key') or '')
        lang = normalize_lang(data.get('language') or '')
        plat = normalize_platform(data.get('platform') or '')
        if data.get('id'):
            existing = CrisisHotline.objects.filter(id=data['id']).first()
            if existing:
                denied = _deny_app_write(request, existing.app_key or '')
                if denied:
                    return denied
            denied = _deny_app_write(request, app)
            if denied:
                return denied
            patch = {}
            for k in ('country', 'label', 'phone_or_url', 'enabled', 'language', 'platform', 'app_key'):
                if k in request.data:
                    if k == 'language':
                        patch[k] = lang
                    elif k == 'platform':
                        patch[k] = plat
                    elif k == 'app_key':
                        patch[k] = app
                    else:
                        patch[k] = data.get(k)
            if patch:
                CrisisHotline.objects.filter(id=data['id']).update(**patch)
            return ApiResponse(data={'id': data['id']})
        denied = _guard_scoped_write(request, target_app=app, require_app_key=True)
        if denied:
            return denied
        r = CrisisHotline.objects.create(
            country=data.get('country') or '',
            label=data['label'],
            phone_or_url=data['phone_or_url'],
            enabled=bool(data.get('enabled', True)),
            app_key=app,
            language=lang,
            platform=plat,
        )
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('快捷开场'))
    @action(detail=False, methods=['get', 'post'], url_path='quick-prompts')
    def quick_prompts(self, request):
        """AI 快捷开场列表或保存/删除。

        业务说明:
            App 侧按语言/平台作用域过滤启用项；此处运营可维护多维度副本。
        """
        if request.method == 'GET':
            qs = AiQuickPrompt.objects.all().order_by('sort_order', 'id')
            qs = _apply_scope_filters(qs, request)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([
                {
                    'id': r.id, 'title': r.title, 'prompt': r.prompt,
                    'sort_order': r.sort_order, 'enabled': r.enabled,
                    'app_key': r.app_key or '',
                    'language': r.language or '', 'platform': r.platform or '',
                }
                for r in page
            ])
        ser = QuickPromptSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        if data.get('delete_id'):
            row = AiQuickPrompt.objects.filter(id=data['delete_id']).first()
            if row:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
                row.delete()
            return ApiResponse(message='deleted')
        app = normalize_app_key(data.get('app_key') or '')
        lang = normalize_lang(data.get('language') or '')
        plat = normalize_platform(data.get('platform') or '')
        if data.get('id'):
            existing = AiQuickPrompt.objects.filter(id=data['id']).first()
            if existing:
                denied = _deny_app_write(request, existing.app_key or '')
                if denied:
                    return denied
            denied = _deny_app_write(request, app)
            if denied:
                return denied
            patch = {}
            for k in ('title', 'prompt', 'sort_order', 'enabled', 'language', 'platform', 'app_key'):
                if k in request.data:
                    if k == 'language':
                        patch[k] = lang
                    elif k == 'platform':
                        patch[k] = plat
                    elif k == 'app_key':
                        patch[k] = app
                    else:
                        patch[k] = data.get(k)
            if patch:
                AiQuickPrompt.objects.filter(id=data['id']).update(**patch)
            return ApiResponse(data={'id': data['id']})
        denied = _guard_scoped_write(request, target_app=app, require_app_key=True)
        if denied:
            return denied
        r = AiQuickPrompt.objects.create(
            title=data['title'],
            prompt=data['prompt'],
            sort_order=int(data.get('sort_order') or 0),
            enabled=bool(data.get('enabled', True)),
            app_key=app,
            language=lang,
            platform=plat,
        )
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('Welcome 欢迎页'))
    @action(detail=False, methods=['get', 'post'], url_path='welcome-screens')
    def welcome_screens(self, request):
        """冷启 Welcome 列表或保存/删除。

        业务说明:
            App 按语言/平台取启用项；version 上调可强制客户端再看。
            表空时 ensure_default_welcome_screens 补一条全局默认。
        """
        ensure_default_welcome_screens()
        if request.method == 'GET':
            qs = WelcomeScreen.objects.all().order_by('-priority', 'sort_order', 'id')
            qs = _apply_scope_filters(qs, request)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([
                {
                    'id': r.id, 'title': r.title, 'subtitle': r.subtitle,
                    'cta_text': r.cta_text, 'image_url': r.image_url,
                    'version': r.version, 'priority': int(r.priority or 0),
                    'sort_order': r.sort_order, 'enabled': r.enabled,
                    'app_key': r.app_key or '',
                    'language': r.language or '', 'platform': r.platform or '',
                }
                for r in page
            ])
        ser = WelcomeScreenSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        if data.get('delete_id'):
            row = WelcomeScreen.objects.filter(id=data['delete_id']).first()
            if row:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
                row.delete()
            return ApiResponse(message='deleted')
        app = normalize_app_key(data.get('app_key') or '')
        lang = normalize_lang(data.get('language') or '')
        plat = normalize_platform(data.get('platform') or '')
        denied = _deny_app_write(request, app)
        if denied:
            return denied
        if data.get('id'):
            existing = WelcomeScreen.objects.filter(id=data['id']).first()
            if existing:
                denied = _deny_app_write(request, existing.app_key or '')
                if denied:
                    return denied
            patch = {}
            for k in (
                'title', 'subtitle', 'cta_text', 'image_url', 'version',
                'priority', 'sort_order', 'enabled', 'language', 'platform', 'app_key',
            ):
                if k in request.data:
                    if k == 'language':
                        patch[k] = lang
                    elif k == 'platform':
                        patch[k] = plat
                    elif k == 'app_key':
                        patch[k] = app
                    else:
                        patch[k] = data.get(k)
            if patch:
                WelcomeScreen.objects.filter(id=data['id']).update(**patch)
            return ApiResponse(data={'id': data['id']})
        denied = _deny_empty_app_create(app)
        if denied:
            return denied
        r = WelcomeScreen.objects.create(
            title=data['title'],
            subtitle=data.get('subtitle') or '',
            cta_text=data.get('cta_text') or 'Continue',
            image_url=data.get('image_url') or '',
            version=int(data.get('version') or 1),
            priority=int(data.get('priority') or 0),
            sort_order=int(data.get('sort_order') or 0),
            enabled=bool(data.get('enabled', True)),
            app_key=app,
            language=lang,
            platform=plat,
        )
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('教育夹层'))
    @action(detail=False, methods=['get', 'post'], url_path='interstitial-screens')
    def interstitial_screens(self, request):
        """教育夹层列表或保存/删除（需 PERM_OPS）。

        业务说明:
            与 Welcome 同构，额外 kind 字段；客户端按 kind + version 门闸。
        """
        denied = self._require_ops(request)
        if denied:
            return denied
        if request.method == 'GET':
            qs = InterstitialScreen.objects.all().order_by('kind', '-priority', 'sort_order', 'id')
            qs = _apply_scope_filters(qs, request)
            kind = (request.query_params.get('kind') or '').strip()
            if kind:
                qs = qs.filter(kind=kind)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([
                {
                    'id': r.id, 'kind': r.kind, 'title': r.title, 'subtitle': r.subtitle,
                    'cta_text': r.cta_text, 'image_url': r.image_url,
                    'version': r.version, 'priority': int(r.priority or 0),
                    'sort_order': r.sort_order, 'enabled': r.enabled,
                    'app_key': r.app_key or '',
                    'language': r.language or '', 'platform': r.platform or '',
                }
                for r in page
            ])
        ser = InterstitialScreenSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        if data.get('delete_id'):
            row = InterstitialScreen.objects.filter(id=data['delete_id']).first()
            if row:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
                row.delete()
            return ApiResponse(message='deleted')
        app = normalize_app_key(data.get('app_key') or '')
        lang = normalize_lang(data.get('language') or '')
        plat = normalize_platform(data.get('platform') or '')
        if data.get('id'):
            existing = InterstitialScreen.objects.filter(id=data['id']).first()
            if existing:
                denied = _deny_app_write(request, existing.app_key or '')
                if denied:
                    return denied
            denied = _deny_app_write(request, app)
            if denied:
                return denied
            patch = {}
            for k in (
                'kind', 'title', 'subtitle', 'cta_text', 'image_url', 'version',
                'priority', 'sort_order', 'enabled', 'language', 'platform', 'app_key',
            ):
                if k in request.data:
                    if k == 'language':
                        patch[k] = lang
                    elif k == 'platform':
                        patch[k] = plat
                    elif k == 'app_key':
                        patch[k] = app
                    else:
                        patch[k] = data.get(k)
            if patch:
                InterstitialScreen.objects.filter(id=data['id']).update(**patch)
            return ApiResponse(data={'id': data['id']})
        denied = _guard_scoped_write(request, target_app=app, require_app_key=True)
        if denied:
            return denied
        r = InterstitialScreen.objects.create(
            kind=data['kind'],
            title=data['title'],
            subtitle=data.get('subtitle') or '',
            cta_text=data.get('cta_text') or 'Continue',
            image_url=data.get('image_url') or '',
            version=int(data.get('version') or 1),
            priority=int(data.get('priority') or 0),
            sort_order=int(data.get('sort_order') or 0),
            enabled=bool(data.get('enabled', True)),
            app_key=app,
            language=lang,
            platform=plat,
        )
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('主题列表'))
    @action(detail=False, methods=['get'], url_path='topics')
    def topics(self, request):
        """主题分页列表（可按 app_key×language×platform 筛选）。

        默认排除软删；trashed=1 只看回收站。
        """
        trashed = (request.query_params.get('trashed') or '').strip().lower() in ('1', 'true', 'yes')
        qs = Topic.objects.filter(is_deleted=trashed).order_by('-sort_weight', '-id')
        qs = _apply_scope_filters(qs, request)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        return paginator.get_paginated_response([
            {
                'id': t.id, 'name': t.name, 'slug': t.slug or '',
                'cover_url': t.cover_url,
                'description': t.description, 'is_vip': t.is_vip,
                'category': (t.category or t.slug or ''),
                'app_key': t.app_key or '',
                'language': t.language or '',
                'platform': t.platform or '',
                'sort_weight': t.sort_weight, 'status': t.status,
            }
            for t in page
        ])

    @extend_schema(summary=_('主题/壁纸简易保存'))
    @action(detail=False, methods=['post'], url_path='topic-save')
    def topic_save(self, request):
        """新建或更新主题。

        业务说明:
            保存时强制 status=1（上架）；下架走软删接口而非改 status。
            slug 由请求传入（类型主题 bible_verses/prayers/quotes/affirmations）；
            未传则保留原值，新建再从 name 生成。
        """
        ser = TopicSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        tid = data.pop('id', None)
        fields = {**data, 'status': 1}
        if tid:
            existing = Topic.objects.filter(id=tid).first()
            if not existing:
                return ApiResponse(message='not found', code=404)
            denied = _deny_app_write(request, existing.app_key or '')
            if denied:
                return denied
            fields['slug'] = _resolve_topic_slug(fields, existing=existing)
            _fill_topic_category(fields, existing=existing)
            if 'app_key' in fields:
                denied = _deny_app_write(request, fields.get('app_key') or '')
                if denied:
                    return denied
            try:
                Topic.objects.filter(id=tid).update(**fields)
            except IntegrityError:
                return ApiResponse(message=_('主题保存冲突，请检查 slug / 名称是否重复'), code=400)
            return ApiResponse(data={'id': tid, 'slug': fields.get('slug') or '', 'category': fields.get('category') or ''})
        denied = _deny_app_write(request, fields.get('app_key') or '')
        if denied:
            return denied
        denied = _deny_empty_app_create(fields.get('app_key') or '')
        if denied:
            return denied
        fields['slug'] = _resolve_topic_slug(fields)
        _fill_topic_category(fields)
        try:
            t = Topic.objects.create(**fields)
        except IntegrityError:
            return ApiResponse(message=_('主题保存冲突，请检查 slug / 名称是否重复'), code=400)
        return ApiResponse(data={'id': t.id, 'slug': t.slug or '', 'category': t.category or ''}, code=201)

    @extend_schema(summary=_('壁纸列表'))
    @action(detail=False, methods=['get'], url_path='wallpapers')
    def wallpapers(self, request):
        """壁纸分页列表（可按 app_key×language×platform、media_type、回收站、关键词筛选）。

        Query:
            q: 唯一标识(id 精确) 或 名称(name_en/es/pt)/主题(theme) 模糊
            is_vip / status（published/disabled 或 1/0）/ theme / visual / usage_min / usage_max / trashed
            theme / visual：tag_index 精确命中，并回落 theme / visual_tags（兼容 {name}）
        """
        trashed = (request.query_params.get('trashed') or '').strip().lower() in ('1', 'true', 'yes')
        qs = Wallpaper.objects.filter(is_deleted=trashed)
        qs = _apply_scope_filters(qs, request)
        media_type = (request.query_params.get('media_type') or '').strip().lower()
        if media_type in ('static', 'video'):
            qs = qs.filter(media_type=media_type)

        q = (request.query_params.get('q') or '').strip()
        if q:
            name_q = (
                Q(name_en__icontains=q)
                | Q(name_es__icontains=q)
                | Q(name_pt__icontains=q)
                | Q(theme__icontains=q)
            )
            if q.isdigit():
                name_q |= Q(pk=int(q))
            qs = qs.filter(name_q)

        vip = (request.query_params.get('is_vip') or '').strip().lower()
        if vip in ('1', 'true', 'yes'):
            qs = qs.filter(is_vip=True)
        elif vip in ('0', 'false', 'no'):
            qs = qs.filter(is_vip=False)

        st = _parse_list_status(request.query_params.get('status'))
        if st is not None:
            qs = qs.filter(status=st)

        theme = (request.query_params.get('theme') or '').strip()
        if theme:
            qs = qs.filter(_typed_tag_q('theme', theme, extra=Q(theme=theme)))

        visual = (
            request.query_params.get('visual')
            or request.query_params.get('visual_tag')
            or ''
        ).strip()
        if visual:
            qs = qs.filter(_typed_tag_q(
                'visual', visual, extra=_json_tag_contains_q('visual_tags', visual),
            ))

        usage_sub = (
            WallpaperUsage.objects.filter(wallpaper_id=OuterRef('pk'))
            .order_by()
            .values('wallpaper_id')
            .annotate(c=Count('id'))
            .values('c')
        )
        qs = qs.annotate(
            usage_count=Coalesce(Subquery(usage_sub, output_field=IntegerField()), Value(0)),
        )

        usage_min = (request.query_params.get('usage_min') or '').strip()
        usage_max = (request.query_params.get('usage_max') or '').strip()
        if usage_min.isdigit():
            qs = qs.filter(usage_count__gte=int(usage_min))
        if usage_max.isdigit():
            qs = qs.filter(usage_count__lte=int(usage_max))

        qs = qs.order_by('-id')
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        ser = WallpaperAdminSerializer(page, many=True)
        return paginator.get_paginated_response(ser.data)

    @extend_schema(summary=_('壁纸保存'))
    @action(detail=False, methods=['post'], url_path='wallpaper-save')
    def wallpaper_save(self, request):
        """新建或更新壁纸。

        业务说明:
            更新只 patch 请求体里出现的键（可只传 id+status 停用/上架）。
            请求里带了图片或视频时，会按类型改写 image_url、video_url、poster_url。
            新建未传 status 时默认上架 1；删除仍走 soft-delete。
        """
        ser = WallpaperSaveSerializer(data=request.data, partial=bool(request.data.get('id')))
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        wid = data.pop('id', None)
        raw = _payload_dict(request)
        fields = _wallpaper_write_fields(request, data)
        if wid:
            existing = Wallpaper.objects.filter(id=wid).first()
            if not existing:
                return ApiResponse(message='not found', code=404)
            _apply_wallpaper_media(fields, raw, data, existing)
            denied = _deny_app_write(request, existing.app_key or '')
            if denied:
                return denied
            if 'app_key' in fields:
                denied = _deny_app_write(request, fields.get('app_key') or '')
                if denied:
                    return denied
            if any(k in fields for k in _WALLPAPER_TAG_FIELDS):
                merged = {
                    'theme': fields.get('theme', existing.theme or ''),
                    'visual_tags': fields.get('visual_tags', existing.visual_tags or []),
                    'scenario_tags': fields.get('scenario_tags', existing.scenario_tags or []),
                    'emotion_tags': fields.get('emotion_tags', existing.emotion_tags or []),
                    'taxonomy': fields.get('taxonomy', existing.taxonomy or {}),
                }
                rebuild_wallpaper_index(merged)
                fields['tag_index'] = merged['tag_index']
            if fields:
                model_fields = {f.name for f in Wallpaper._meta.fields}
                for key, value in fields.items():
                    if key in model_fields:
                        setattr(existing, key, value)
                existing.save()
            existing.refresh_from_db()
            # 自动创建缺失的隐性标签
            ensure_tag_exists('theme', fields.get('theme', existing.theme) or '')
            for _name in (fields.get('visual_tags', existing.visual_tags) or []):
                ensure_tag_exists('visual', _name)
            for _name in (fields.get('scenario_tags', existing.scenario_tags) or []):
                ensure_tag_exists('scenario', _name)
            for _name in (fields.get('emotion_tags', existing.emotion_tags) or []):
                ensure_tag_exists('emotion', _name)
            for _type_key, _names in (fields.get('taxonomy', existing.taxonomy) or {}).items():
                for _name in (_names or []):
                    ensure_tag_exists(_type_key, _name)
            try:
                payload = WallpaperAdminSerializer(existing).data
            except Exception:
                payload = {'id': existing.id}
            return ApiResponse(data=payload)
        fields.setdefault('status', 1)
        _apply_wallpaper_media(fields, raw, data, None)
        denied = _deny_app_write(request, fields.get('app_key') or '')
        if denied:
            return denied
        denied = _deny_empty_app_create(fields.get('app_key') or '')
        if denied:
            return denied
        rebuild_wallpaper_index(fields)
        model_fields = {f.name for f in Wallpaper._meta.fields}
        create_kwargs = {k: v for k, v in fields.items() if k in model_fields}
        w = Wallpaper.objects.create(**create_kwargs)
        # 自动创建缺失的隐性标签
        ensure_tag_exists('theme', fields.get('theme', ''))
        for _name in (fields.get('visual_tags') or []):
            ensure_tag_exists('visual', _name)
        for _name in (fields.get('scenario_tags') or []):
            ensure_tag_exists('scenario', _name)
        for _name in (fields.get('emotion_tags') or []):
            ensure_tag_exists('emotion', _name)
        for _type_key, _names in (fields.get('taxonomy') or {}).items():
            for _name in (_names or []):
                ensure_tag_exists(_type_key, _name)
        try:
            payload = WallpaperAdminSerializer(w).data
        except Exception:
            payload = {'id': w.id}
        return ApiResponse(data=payload, code=201)

    @extend_schema(summary=_('软删除主题'))
    @action(detail=False, methods=['post'], url_path='topic-delete')
    def topic_delete(self, request):
        """软删除主题（置 is_deleted，语录 topic_id 仍保留历史归属）。"""
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        row = Topic.objects.filter(id=ser.validated_data['id']).first()
        if row:
            denied = _deny_app_write(request, row.app_key or '')
            if denied:
                return denied
            Topic.objects.filter(id=row.id).update(is_deleted=True)
        return ApiResponse(message='ok')

    @extend_schema(summary=_('从回收站恢复主题'))
    @action(detail=False, methods=['post'], url_path='topic-restore')
    def topic_restore(self, request):
        """恢复软删除主题（is_deleted=False）。"""
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        row = Topic.objects.filter(id=ser.validated_data['id']).first()
        if not row:
            return ApiResponse(message='not found', code=404)
        denied = _deny_app_write(request, row.app_key or '')
        if denied:
            return denied
        Topic.objects.filter(id=row.id).update(is_deleted=False)
        return ApiResponse(message='ok', data={'id': row.id})

    @extend_schema(summary=_('软删除壁纸（支持批量）'))
    @action(detail=False, methods=['post'], url_path='wallpaper-delete')
    def wallpaper_delete(self, request):
        """软删除壁纸。

        业务说明:
            置 is_deleted=True，不物理删除（用户 wallpaper_id 可能仍指向旧 ID，App 侧需容错）。
            支持 {"id": 1} 单删或 {"ids": [1,2,3]} 批量删。
        """
        ser = BatchIdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        ids = ser.get_ids()
        rows = Wallpaper.objects.filter(id__in=ids)
        if rows.exists():
            for row in rows:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
            Wallpaper.objects.filter(id__in=[r.id for r in rows]).update(is_deleted=True)
        return ApiResponse(message='ok', data={'deleted': len(rows)})

    @extend_schema(summary=_('从回收站恢复壁纸'))
    @action(detail=False, methods=['post'], url_path='wallpaper-restore')
    def wallpaper_restore(self, request):
        """恢复软删除壁纸（is_deleted=False）。"""
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        row = Wallpaper.objects.filter(id=ser.validated_data['id']).first()
        if not row:
            return ApiResponse(message='not found', code=404)
        denied = _deny_app_write(request, row.app_key or '')
        if denied:
            return denied
        Wallpaper.objects.filter(id=row.id).update(is_deleted=False)
        return ApiResponse(message='ok', data={'id': row.id})

    # ---------- Selah 3.0 Quiz / App icons ----------

    @extend_schema(summary=_('测验列表'))
    @action(detail=False, methods=['get'], url_path='quizzes')
    def quizzes(self, request):
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        qs = Quiz.objects.filter(is_deleted=False).order_by('sort_order', 'id')
        qs = _apply_scope_filters(qs, request)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        return paginator.get_paginated_response([
            {
                'id': r.id, 'title': r.title, 'subtitle': r.subtitle,
                'cover_url': r.cover_url, 'sort_order': r.sort_order,
                'enabled': r.enabled, 'is_vip': r.is_vip,
                'app_key': r.app_key or '', 'language': r.language or '',
                'platform': r.platform or '',
                'question_count': QuizQuestion.objects.filter(quiz_id=r.id, enabled=True).count(),
            }
            for r in page
        ])

    @extend_schema(summary=_('测验保存'))
    @action(detail=False, methods=['post'], url_path='quiz-save')
    def quiz_save(self, request):
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        ser = QuizSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = ser.validated_data
        qid = data.pop('id', None)
        if qid:
            existing = Quiz.objects.filter(id=qid).first()
            if not existing:
                return ApiResponse(message='not found', code=404)
            denied = _deny_app_write(request, existing.app_key or '')
            if denied:
                return denied
            if 'app_key' in data:
                denied = _deny_app_write(request, data.get('app_key') or '')
                if denied:
                    return denied
            Quiz.objects.filter(id=qid).update(**data)
            return ApiResponse(data={'id': qid})
        denied = _guard_scoped_write(
            request, target_app=data.get('app_key') or '', require_app_key=True,
        )
        if denied:
            return denied
        row = Quiz.objects.create(**data)
        return ApiResponse(data={'id': row.id}, code=201)

    @extend_schema(summary=_('测验软删'))
    @action(detail=False, methods=['post'], url_path='quiz-delete')
    def quiz_delete(self, request):
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        row = Quiz.objects.filter(id=ser.validated_data['id']).first()
        if row:
            denied = _deny_app_write(request, row.app_key or '')
            if denied:
                return denied
            Quiz.objects.filter(id=row.id).update(is_deleted=True)
        return ApiResponse(message='ok')

    @extend_schema(summary=_('测验题目列表'))
    @action(detail=False, methods=['get'], url_path='quiz-questions')
    def quiz_questions(self, request):
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        try:
            quiz_id = int(request.query_params.get('quiz_id') or 0)
        except (TypeError, ValueError):
            quiz_id = 0
        qs = QuizQuestion.objects.all().order_by('sort_order', 'id')
        if quiz_id:
            qs = qs.filter(quiz_id=quiz_id)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        return paginator.get_paginated_response([
            {
                'id': r.id, 'quiz_id': r.quiz_id, 'stem': r.stem,
                'options': r.options or [], 'correct_key': r.correct_key,
                'explain': r.explain, 'sort_order': r.sort_order, 'enabled': r.enabled,
            }
            for r in page
        ])

    @extend_schema(summary=_('测验题目保存'))
    @action(detail=False, methods=['post'], url_path='quiz-question-save')
    def quiz_question_save(self, request):
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        ser = QuizQuestionSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = ser.validated_data
        qid = data.pop('id', None)
        quiz_id = data.get('quiz_id')
        parent = Quiz.objects.filter(id=quiz_id).first() if quiz_id else None
        if qid:
            existing = QuizQuestion.objects.filter(id=qid).first()
            if not existing:
                return ApiResponse(message='not found', code=404)
            parent = Quiz.objects.filter(id=existing.quiz_id).first()
            if parent:
                denied = _deny_app_write(request, parent.app_key or '')
                if denied:
                    return denied
            QuizQuestion.objects.filter(id=qid).update(**data)
            return ApiResponse(data={'id': qid})
        if not parent:
            return ApiResponse(message='quiz_id required', code=400)
        denied = _deny_app_write(request, parent.app_key or '')
        if denied:
            return denied
        row = QuizQuestion.objects.create(**data)
        return ApiResponse(data={'id': row.id}, code=201)

    @extend_schema(summary=_('测验题目删除'))
    @action(detail=False, methods=['post'], url_path='quiz-question-delete')
    def quiz_question_delete(self, request):
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        existing = QuizQuestion.objects.filter(id=ser.validated_data['id']).first()
        if existing:
            parent = Quiz.objects.filter(id=existing.quiz_id).first()
            if parent:
                denied = _deny_app_write(request, parent.app_key or '')
                if denied:
                    return denied
            existing.delete()
        return ApiResponse(message='ok')

    @extend_schema(summary=_('App图标列表'))
    @action(detail=False, methods=['get'], url_path='app-icons')
    def app_icons_admin(self, request):
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        qs = AppIcon.objects.filter(is_deleted=False).order_by('sort_order', 'id')
        qs = _apply_scope_filters(qs, request)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        return paginator.get_paginated_response([
            {
                'id': r.id, 'icon_key': r.icon_key, 'title': r.title,
                'preview_url': r.preview_url, 'is_vip': r.is_vip,
                'sort_order': r.sort_order, 'enabled': r.enabled,
                'app_key': r.app_key or '', 'language': r.language or '',
                'platform': r.platform or '',
            }
            for r in page
        ])

    @extend_schema(summary=_('App图标保存'))
    @action(detail=False, methods=['post'], url_path='app-icon-save')
    def app_icon_save(self, request):
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        ser = AppIconSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = ser.validated_data
        iid = data.pop('id', None)
        if iid:
            existing = AppIcon.objects.filter(id=iid).first()
            if not existing:
                return ApiResponse(message='not found', code=404)
            denied = _deny_app_write(request, existing.app_key or '')
            if denied:
                return denied
            if 'app_key' in data:
                denied = _deny_app_write(request, data.get('app_key') or '')
                if denied:
                    return denied
            AppIcon.objects.filter(id=iid).update(**data)
            return ApiResponse(data={'id': iid})
        denied = _guard_scoped_write(
            request, target_app=data.get('app_key') or '', require_app_key=True,
        )
        if denied:
            return denied
        row = AppIcon.objects.create(**data)
        return ApiResponse(data={'id': row.id}, code=201)

    @extend_schema(summary=_('App图标软删'))
    @action(detail=False, methods=['post'], url_path='app-icon-delete')
    def app_icon_delete(self, request):
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        row = AppIcon.objects.filter(id=ser.validated_data['id']).first()
        if row:
            denied = _deny_app_write(request, row.app_key or '')
            if denied:
                return denied
            AppIcon.objects.filter(id=row.id).update(is_deleted=True)
        return ApiResponse(message='ok')

    # ---------- 用户侧（App）用户管理 ----------

    @extend_schema(summary=_('App 用户列表'))
    @action(detail=False, methods=['get'], url_path='app-users')
    def app_users(self, request):
        """App 用户分页检索。

        业务说明:
            仅 role=user；q 模糊匹配用户名/邮箱/设备/称呼。
            status 仅接受 0/1；vip_level / register_country 精确过滤。
            需 PERM_APP_USERS。
        """
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message=_('无权限'), code=403)
        qs = User.objects.filter(role='user').order_by('-id')
        allowed_apps = staff_managed_apps(request.user)
        if allowed_apps:
            qs = qs.filter(app_key__in=allowed_apps)
        q = (request.query_params.get('q') or '').strip()
        if q:
            qs = qs.filter(
                Q(username__icontains=q)
                | Q(email__icontains=q)
                | Q(device_id__icontains=q)
                | Q(display_name__icontains=q)
            )
        status = request.query_params.get('status')
        if status in ('0', '1'):
            qs = qs.filter(status=int(status))
        vip = (request.query_params.get('vip_level') or '').strip()
        if vip:
            qs = qs.filter(vip_level=vip)
        country = (request.query_params.get('register_country') or '').strip()
        if country:
            qs = qs.filter(register_country__iexact=country)
        app_key = normalize_app_key(request.query_params.get('app_key') or '')
        if app_key:
            if allowed_apps and app_key not in allowed_apps:
                qs = qs.none()
            else:
                qs = qs.filter(app_key=app_key)
        lang = normalize_lang(
            request.query_params.get('system_lang')
            or request.query_params.get('language')
            or ''
        )
        if lang:
            qs = qs.filter(system_lang=lang)
        plat = normalize_platform(request.query_params.get('platform') or '')
        if plat:
            qs = qs.filter(platform=plat)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        return paginator.get_paginated_response([_serialize_app_user(u) for u in page])

    @extend_schema(summary=_('App 用户详情'))
    @action(detail=False, methods=['get'], url_path='app-user-detail')
    def app_user_detail(self, request):
        """App 用户详情。"""
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message=_('无权限'), code=403)
        uid = request.query_params.get('id')
        if not uid:
            return ApiResponse(message='id required', code=400)
        try:
            user = User.objects.get(id=uid, role='user')
        except User.DoesNotExist:
            return ApiResponse(message=_('用户不存在'), code=404)
        return ApiResponse(data=_serialize_app_user(user))

    @extend_schema(summary=_('App 用户保存'))
    @action(detail=False, methods=['post'], url_path='app-user-save')
    def app_user_save(self, request):
        """更新 App 用户状态/VIP/额度等。

        业务说明:
            仅 patch 请求体出现的白名单字段（含 free_quota / vip_daily_remaining 人工调额）。
            不可改 role，防止把 App 用户提成运营。
        """
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message=_('无权限'), code=403)
        ser = AppUserSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = ser.validated_data
        try:
            user = User.objects.get(id=data['id'], role='user')
        except User.DoesNotExist:
            return ApiResponse(message=_('用户不存在'), code=404)
        patch = {}
        for key in (
            'status', 'display_name', 'vip_level', 'vip_expire_at', 'is_trial',
            'free_quota', 'vip_daily_remaining', 'marketing_opt_in',
        ):
            if key in request.data:
                patch[key] = data.get(key)
        if patch:
            for k, v in patch.items():
                setattr(user, k, v)
            user.save(update_fields=list(patch.keys()) + ['updated_at'])
        return ApiResponse(data=_serialize_app_user(user))

    @extend_schema(summary=_('App 用户启用/禁用'))
    @action(detail=False, methods=['post'], url_path='app-user-toggle-status')
    def app_user_toggle_status(self, request):
        """启用或禁用 App 用户（1↔0 翻转；禁用后登录/鉴权侧应拒绝）。"""
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message=_('无权限'), code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        try:
            user = User.objects.get(id=ser.validated_data['id'], role='user')
        except User.DoesNotExist:
            return ApiResponse(message=_('用户不存在'), code=404)
        user.status = 0 if user.status == 1 else 1
        user.save(update_fields=['status', 'updated_at'])
        return ApiResponse(data={'id': user.id, 'status': user.status})

    # ---------- 后台工作人员管理 ----------

    @extend_schema(summary=_('后台用户列表'))
    @action(detail=False, methods=['get'], url_path='staff-users')
    def staff_users(self, request):
        """运营账号列表（role∈STAFF_ROLES；可按角色/状态/关键词筛）。"""
        if not has_perm(request.user, PERM_STAFF):
            return ApiResponse(message=_('无权限'), code=403)
        qs = User.objects.filter(role__in=STAFF_ROLES).order_by('-id')
        q = (request.query_params.get('q') or '').strip()
        if q:
            qs = qs.filter(Q(username__icontains=q) | Q(email__icontains=q))
        role = (request.query_params.get('role') or '').strip()
        if role in STAFF_ROLES:
            qs = qs.filter(role=role)
        status = request.query_params.get('status')
        if status in ('0', '1'):
            qs = qs.filter(status=int(status))
        app_key = normalize_app_key(request.query_params.get('app_key') or '')
        if app_key:
            # 指定 App：含该 App，或 app_keys 为空（全部 App）
            qs = qs.filter(Q(app_keys__contains=[app_key]) | Q(app_keys=[]) | Q(app_keys__isnull=True))
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        return paginator.get_paginated_response([_serialize_staff(u) for u in page])

    @extend_schema(summary=_('后台用户保存'))
    @action(detail=False, methods=['post'], url_path='staff-user-save')
    def staff_user_save(self, request):
        """新建或更新运营账号。

        业务说明（RBAC 护栏）:
            - 仅 super_admin 可分配/创建 super_admin。
            - 非超管不可改已有超管账号。
            - 不可禁用自己、不可改自己的角色。
            - 密码明文入参，User.save 内 bcrypt 哈希。
        """
        if not has_perm(request.user, PERM_STAFF):
            return ApiResponse(message=_('无权限'), code=403)
        ser = StaffUserSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        actor = request.user
        sid = data.pop('id', None)
        password = data.pop('password', None) or None
        target_role = data.get('role')

        if target_role == 'super_admin' and actor.role != 'super_admin':
            return ApiResponse(message=_('仅超级管理员可分配该角色'), code=403)

        if sid:
            try:
                user = User.objects.get(id=sid, role__in=STAFF_ROLES)
            except User.DoesNotExist:
                return ApiResponse(message=_('账号不存在'), code=404)
            if user.role == 'super_admin' and actor.role != 'super_admin':
                return ApiResponse(message=_('无权修改超级管理员'), code=403)
            # 防把自己锁死在运营台外
            if user.id == actor.id and 'status' in request.data and data.get('status') == 0:
                return ApiResponse(message=_('不能禁用自己'), code=400)
            if user.id == actor.id and target_role and target_role != user.role:
                return ApiResponse(message=_('不能修改自己的角色'), code=400)
            if 'username' in request.data and data.get('username'):
                if User.objects.filter(username=data['username']).exclude(id=user.id).exists():
                    return ApiResponse(message=_('用户名已存在'), code=400)
                user.username = data['username']
            if 'email' in request.data:
                email = data.get('email') or None
                app_k = getattr(user, 'app_key', '') or ''
                if email and User.objects.filter(email=email, app_key=app_k).exclude(id=user.id).exists():
                    return ApiResponse(message=_('邮箱已存在'), code=400)
                user.email = email
            if target_role:
                user.role = target_role
            if 'status' in request.data:
                user.status = data['status']
            if 'app_keys' in request.data:
                user.app_keys = data.get('app_keys') or []
            if password:
                user.password = password  # User.save 会哈希
            user.save()
            return ApiResponse(data=_serialize_staff(user))

        # 新建：默认 role=operator；login_type=password
        username = data.get('username')
        if User.objects.filter(username=username).exists():
            return ApiResponse(message=_('用户名已存在'), code=400)
        email = data.get('email') or None
        # 新建运营账号默认 app_key=''；唯一约束为 (email, app_key)
        if email and User.objects.filter(email=email, app_key='').exists():
            return ApiResponse(message=_('邮箱已存在'), code=400)
        role = target_role or 'operator'
        if role == 'super_admin' and actor.role != 'super_admin':
            return ApiResponse(message=_('仅超级管理员可创建超级管理员'), code=403)
        user = User(
            username=username,
            email=email,
            password=password,
            role=role,
            status=data.get('status', 1),
            login_type='password',
            app_keys=data.get('app_keys') or [],
        )
        user.save()
        return ApiResponse(data=_serialize_staff(user), code=201)

    @extend_schema(summary=_('后台用户启用/禁用'))
    @action(detail=False, methods=['post'], url_path='staff-user-toggle-status')
    def staff_user_toggle_status(self, request):
        """启用或禁用运营账号。

        业务说明:
            禁止禁用自己；非超管不可动超管账号。
        """
        if not has_perm(request.user, PERM_STAFF):
            return ApiResponse(message=_('无权限'), code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        try:
            user = User.objects.get(id=ser.validated_data['id'], role__in=STAFF_ROLES)
        except User.DoesNotExist:
            return ApiResponse(message=_('账号不存在'), code=404)
        if user.id == request.user.id:
            return ApiResponse(message=_('不能禁用自己'), code=400)
        if user.role == 'super_admin' and request.user.role != 'super_admin':
            return ApiResponse(message=_('无权操作超级管理员'), code=403)
        user.status = 0 if user.status == 1 else 1
        user.save(update_fields=['status', 'updated_at'])
        return ApiResponse(data={'id': user.id, 'status': user.status})

    # ---------- 权限管理 ----------

    @extend_schema(summary=_('角色权限矩阵'))
    @action(detail=False, methods=['get', 'post'], url_path='permissions')
    def permissions(self, request):
        """查询或保存角色权限矩阵。

        业务说明:
            GET：有 PERM_PERMISSIONS 返回完整矩阵；否则只回本人角色权限（供侧栏菜单）。
            POST：需 PERM_PERMISSIONS；roles 为 {role: [perm_keys...]}，落库后立即生效。
        """
        if request.method == 'GET':
            # 登录后拉取菜单权限：任意管理员可看自己；完整矩阵需 permissions 权限
            if has_perm(request.user, PERM_PERMISSIONS):
                return ApiResponse(data={
                    'catalog': permission_catalog(),
                    'roles': get_role_permissions_map(),
                    'role_labels': ROLE_LABELS,
                })
            return ApiResponse(data={
                'catalog': permission_catalog(),
                'roles': {request.user.role: permissions_for_role(request.user.role)},
                'role_labels': ROLE_LABELS,
                'mine': permissions_for_role(request.user.role),
            })

        if not has_perm(request.user, PERM_PERMISSIONS):
            return ApiResponse(message=_('无权限'), code=403)
        ser = RolePermissionsSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        # 持久化角色→权限映射（影响后续 has_perm）
        saved = save_role_permissions(ser.validated_data['roles'])
        return ApiResponse(data={
            'catalog': permission_catalog(),
            'roles': saved,
            'role_labels': ROLE_LABELS,
        })

    # ---------- 会员商品 ----------

    @extend_schema(summary=_('会员商品列表'))
    @action(detail=False, methods=['get'], url_path='membership-products')
    def membership_products(self, request):
        """会员商品列表。

        业务说明:
            ensure_default_products 保证空库有种子商品；可按平台/套餐/启用过滤。
            meta.trial_days 为当前 query 维度下的默认试用天数（商品可覆盖）；
            可传 app_key/language/platform，空=全局。
        """
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        ensure_default_products()
        qs = MembershipProduct.objects.all().order_by('-sort_weight', 'app_key', 'plan_type', 'platform', 'id')
        qs = apply_staff_app_list_filter(qs, request.user, field='app_key')
        app_q = request.query_params.get('app_key')
        if app_q is not None:
            ak = normalize_app_key(app_q)
            allowed = staff_managed_apps(request.user)
            if allowed and ak and ak not in allowed:
                qs = qs.none()
            else:
                qs = qs.filter(app_key=ak)
        platform = (request.query_params.get('platform') or '').strip()
        if platform in ('ios', 'android'):
            qs = qs.filter(platform=platform)
        plan_type = (request.query_params.get('plan_type') or '').strip()
        if plan_type in ('monthly', 'yearly'):
            qs = qs.filter(plan_type=plan_type)
        enabled = request.query_params.get('enabled')
        if enabled in ('0', '1', 'true', 'false'):
            qs = qs.filter(enabled=enabled in ('1', 'true'))
        from tools.selah_quota import get_config
        app = normalize_app_key(request.query_params.get('app_key') or '')
        lang = normalize_lang(request.query_params.get('language') or '')
        plat = normalize_platform(request.query_params.get('platform') or '')
        return ApiResponse(
            data=[serialize_product(p) for p in qs],
            meta={
                'trial_days': int(get_config('trial_days', 3, language=lang, platform=plat, app_key=app)),
                'app_key': app,
                'language': lang,
                'platform': plat,
            },
        )

    @extend_schema(summary=_('会员商品保存'))
    @action(detail=False, methods=['post'], url_path='membership-product-save')
    def membership_product_save(self, request):
        """新建或更新会员商品并同步 IAP 配置键。

        业务说明:
            (platform, store_product_id) 唯一；保存后 sync_iap_config_keys 回写 AppConfig 里的
            iap_*_product_id，供 App 公开配置读取。
            新建 yearly 默认 duration_days=365，否则 30。
        """
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        ser = MembershipProductSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        pid = data.pop('id', None)
        store_id = (data.get('store_product_id') or '').strip() if 'store_product_id' in data else None

        if pid:
            try:
                row = MembershipProduct.objects.get(id=pid)
            except MembershipProduct.DoesNotExist:
                return ApiResponse(message=_('商品不存在'), code=404)
            target_app = normalize_app_key(data.get('app_key')) if 'app_key' in request.data else (row.app_key or '')
            denied = _deny_app_write(request, target_app)
            if denied:
                return denied
            denied = _deny_app_write(request, row.app_key or '')
            if denied:
                return denied
            # 改商店 ID 时校验同平台唯一
            if store_id is not None:
                platform = data.get('platform') or row.platform
                if MembershipProduct.objects.filter(platform=platform, store_product_id=store_id).exclude(id=pid).exists():
                    return ApiResponse(message=_('该平台商店商品ID已存在'), code=400)
            for key in (
                'app_key', 'plan_type', 'platform', 'store_product_id', 'title', 'price_display',
                'price_display_i18n', 'price_cents', 'currency',
                'duration_days', 'trial_days', 'enabled', 'sort_weight', 'remark',
            ):
                if key in request.data:
                    if key == 'app_key':
                        setattr(row, key, normalize_app_key(data.get(key) or ''))
                    else:
                        setattr(row, key, data.get(key))
            row.save()
            sync_iap_config_keys(app_key=row.app_key or '')
            return ApiResponse(data=serialize_product(row))

        app = normalize_app_key(data.get('app_key') or '')
        denied = _deny_app_write(request, app)
        if denied:
            return denied
        if not app:
            return ApiResponse(message=_('请指定 App（app_key）'), code=400)
        platform = data['platform']
        if MembershipProduct.objects.filter(platform=platform, store_product_id=store_id).exists():
            return ApiResponse(message=_('该平台商店商品ID已存在'), code=400)
        row = MembershipProduct.objects.create(
            app_key=app,
            plan_type=data['plan_type'],
            platform=platform,
            store_product_id=store_id,
            title=data.get('title') or '',
            price_display=data.get('price_display') or '',
            price_display_i18n=data.get('price_display_i18n') or {},
            price_cents=int(data.get('price_cents') or 0),
            currency=(data.get('currency') or 'USD')[:8],
            duration_days=int(data.get('duration_days') or (365 if data['plan_type'] == 'yearly' else 30)),
            trial_days=data.get('trial_days'),
            enabled=bool(data.get('enabled', True)),
            sort_weight=int(data.get('sort_weight') or 0),
            remark=data.get('remark') or f'pack:{app}',
        )
        sync_iap_config_keys(app_key=app)
        return ApiResponse(data=serialize_product(row), code=201)

    @extend_schema(summary=_('会员商品删除'))
    @action(detail=False, methods=['post'], url_path='membership-product-delete')
    def membership_product_delete(self, request):
        """删除会员商品（硬删；随后同步该 App 的 IAP 配置键）。"""
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        row = MembershipProduct.objects.filter(id=ser.validated_data['id']).first()
        if not row:
            return ApiResponse(message=_('商品不存在'), code=404)
        denied = _deny_app_write(request, row.app_key or '')
        if denied:
            return denied
        app = row.app_key or ''
        row.delete()
        sync_iap_config_keys(app_key=app)
        return ApiResponse(message='ok')

    # ---------- 订单（IAP 收据） ----------

    def _serialize_order(self, row, user_map=None):
        """订单行序列化（附用户摘要与 App）。"""
        user = (user_map or {}).get(row.user_id)
        app = getattr(row, 'app_key', '') or ''
        if not app and user:
            app = getattr(user, 'app_key', '') or ''
        return {
            'id': row.id,
            'user_id': row.user_id,
            'username': (user.username if user else '') or '',
            'display_name': (user.display_name if user else '') or '',
            'device_id': (user.device_id if user else '') or '',
            'app_key': app,
            'platform': row.platform,
            'product_id': row.product_id,
            'transaction_id': row.transaction_id,
            'original_transaction_id': row.original_transaction_id or '',
            'is_trial': row.is_trial,
            'amount_cents': int(getattr(row, 'amount_cents', 0) or 0),
            'currency': getattr(row, 'currency', None) or 'USD',
            'expire_at': row.expire_at.isoformat() if row.expire_at else None,
            'status': row.status,
            'created_at': row.created_at.isoformat() if row.created_at else None,
            'updated_at': row.updated_at.isoformat() if row.updated_at else None,
        }

    @extend_schema(summary=_('订单列表'))
    @action(detail=False, methods=['get'], url_path='orders')
    def orders(self, request):
        """IAP 订单分页列表。

        业务说明:
            q 同时搜交易号/商品/用户摘要（用户 id 预取上限 500）。
            app_key 按收据 app_key 过滤；空收据回退匹配用户 app_key。
            列表前懒更新：active 且 expire_at 已过 → status=expired（非定时任务）。
        """
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        qs = IapReceipt.objects.all().order_by('-id')
        q = (request.query_params.get('q') or '').strip()
        if q:
            user_ids = list(
                User.objects.filter(
                    Q(username__icontains=q)
                    | Q(device_id__icontains=q)
                    | Q(display_name__icontains=q)
                    | Q(email__icontains=q)
                ).values_list('id', flat=True)[:500]
            )
            q_filter = (
                Q(transaction_id__icontains=q)
                | Q(product_id__icontains=q)
                | Q(original_transaction_id__icontains=q)
                | Q(user_id__in=user_ids)
            )
            if q.isdigit():
                q_filter = q_filter | Q(user_id=int(q)) | Q(id=int(q))
            qs = qs.filter(q_filter)
        platform = (request.query_params.get('platform') or '').strip()
        if platform in ('ios', 'android'):
            qs = qs.filter(platform=platform)
        status = (request.query_params.get('status') or '').strip()
        if status:
            qs = qs.filter(status=status)
        trial = request.query_params.get('is_trial')
        if trial in ('0', '1', 'true', 'false'):
            qs = qs.filter(is_trial=trial in ('1', 'true'))
        app_key = normalize_app_key(request.query_params.get('app_key') or '')
        if app_key:
            app_uids = list(User.objects.filter(app_key=app_key).values_list('id', flat=True))
            qs = qs.filter(Q(app_key=app_key) | (Q(app_key='') & Q(user_id__in=app_uids)))
        # 语言筛用户 system_lang（收据本身无 language 字段）
        lang = normalize_lang(request.query_params.get('language') or '')
        if lang:
            lang_uids = list(User.objects.filter(system_lang=lang).values_list('id', flat=True))
            qs = qs.filter(user_id__in=lang_uids)

        # 过期自动标记（仅列表展示侧懒更新）
        now = dj_tz.now()
        IapReceipt.objects.filter(status='active', expire_at__lt=now).update(status='expired')

        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        uids = {r.user_id for r in page}
        user_map = {u.id: u for u in User.objects.filter(id__in=uids)}
        return paginator.get_paginated_response([self._serialize_order(r, user_map) for r in page])

    @extend_schema(summary=_('订单详情'))
    @action(detail=False, methods=['get'], url_path='order-detail')
    def order_detail(self, request):
        """IAP 订单详情。"""
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        oid = request.query_params.get('id')
        if not oid:
            return ApiResponse(message='id required', code=400)
        try:
            row = IapReceipt.objects.get(id=oid)
        except IapReceipt.DoesNotExist:
            return ApiResponse(message=_('订单不存在'), code=404)
        user = User.objects.filter(id=row.user_id).first()
        data = self._serialize_order(row, {row.user_id: user} if user else {})
        data['receipt_data'] = (row.receipt_data or '')[:500]
        return ApiResponse(data=data)

    @extend_schema(summary=_('订单状态更新'))
    @action(detail=False, methods=['post'], url_path='order-status')
    def order_status(self, request):
        """更新订单状态（退款/吊销等）。

        业务说明:
            refunded/revoked 时：若用户无其它未过期 active 收据，则降级为 free 并清 is_trial。
            仍有其它有效收据则保留 VIP（多笔订阅场景）。
        """
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        ser = OrderStatusSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        try:
            row = IapReceipt.objects.get(id=ser.validated_data['id'])
        except IapReceipt.DoesNotExist:
            return ApiResponse(message=_('订单不存在'), code=404)
        row.status = ser.validated_data['status']
        row.save(update_fields=['status', 'updated_at'])
        # 退款/撤销时降级用户 VIP（仅当该订单仍是当前有效凭证之一）
        if row.status in ('refunded', 'revoked'):
            user = User.objects.filter(id=row.user_id).first()
            if user and user.is_vip_active:
                still = IapReceipt.objects.filter(
                    user_id=user.id, status='active',
                ).exclude(id=row.id).filter(
                    Q(expire_at__isnull=True) | Q(expire_at__gt=dj_tz.now())
                ).exists()
                if not still:
                    user.vip_level = User.VIP_FREE
                    user.is_trial = False
                    user.save(update_fields=['vip_level', 'is_trial', 'updated_at'])
        user = User.objects.filter(id=row.user_id).first()
        return ApiResponse(data=self._serialize_order(row, {row.user_id: user} if user else {}))

    # ---------- 内容举报 / AI 安全审计 ----------

    def _serialize_report(self, row, user_map, quote_map):
        """举报行序列化（附用户与语录摘要）。"""
        u = user_map.get(row.user_id)
        q = quote_map.get(row.content_id)
        return {
            'id': row.id,
            'user_id': row.user_id,
            'display_name': (u.display_name if u else '') or (u.username if u else '') or '',
            'device_id': u.device_id if u else '',
            'app_key': (getattr(u, 'app_key', '') if u else '') or (getattr(q, 'app_key', '') if q else '') or '',
            'content_id': row.content_id,
            'quote_content': (q.content[:200] if q else ''),
            'quote_deleted': bool(q.is_deleted) if q else True,
            'reason': row.reason,
            'detail': row.detail or '',
            'status': row.status,
            'created_at': row.created_at.isoformat() if row.created_at else None,
            'updated_at': row.updated_at.isoformat() if getattr(row, 'updated_at', None) else None,
        }

    @extend_schema(summary=_('内容举报列表'))
    @action(detail=False, methods=['get'], url_path='content-reports')
    def content_reports(self, request):
        """内容举报列表。

        业务说明:
            status∈pending/handled/dismissed；q 可搜语录正文/用户/详情，数字则兼匹配 id。
            app_key 按举报用户过滤。需 PERM_CONTENT。
        """
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        qs = ContentReport.objects.all().order_by('-id')
        status = (request.query_params.get('status') or '').strip()
        if status in ('pending', 'handled', 'dismissed'):
            qs = qs.filter(status=status)
        reason = (request.query_params.get('reason') or '').strip()
        if reason:
            qs = qs.filter(reason=reason)
        app_uids = _user_ids_for_scope_query(request)
        if app_uids is not None:
            qs = qs.filter(user_id__in=app_uids)
        q = (request.query_params.get('q') or '').strip()
        if q:
            quote_ids = list(
                Quote.objects.filter(content__icontains=q).values_list('id', flat=True)[:500]
            )
            user_ids = list(
                User.objects.filter(
                    Q(display_name__icontains=q) | Q(device_id__icontains=q) | Q(username__icontains=q)
                ).values_list('id', flat=True)[:500]
            )
            q_filter = Q(content_id__in=quote_ids) | Q(user_id__in=user_ids) | Q(detail__icontains=q)
            if q.isdigit():
                q_filter = q_filter | Q(id=int(q)) | Q(content_id=int(q)) | Q(user_id=int(q))
            qs = qs.filter(q_filter)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        uids = {r.user_id for r in page}
        cids = {r.content_id for r in page}
        user_map = {u.id: u for u in User.objects.filter(id__in=uids)}
        quote_map = {qq.id: qq for qq in Quote.objects.filter(id__in=cids)}
        return paginator.get_paginated_response(
            [self._serialize_report(r, user_map, quote_map) for r in page]
        )

    @extend_schema(summary=_('内容举报处理'))
    @action(detail=False, methods=['post'], url_path='content-report-status')
    def content_report_status(self, request):
        """处理举报（可顺带软删语录）。

        业务说明:
            soft_delete_quote=True 时同步软删被举报语录（与 quote-delete 同效果）。
        """
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        ser = ReportStatusSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        try:
            row = ContentReport.objects.get(id=ser.validated_data['id'])
        except ContentReport.DoesNotExist:
            return ApiResponse(message=_('举报不存在'), code=404)
        row.status = ser.validated_data['status']
        row.save(update_fields=['status', 'updated_at'])
        if ser.validated_data.get('soft_delete_quote'):
            Quote.objects.filter(id=row.content_id).update(is_deleted=True)
        user = User.objects.filter(id=row.user_id).first()
        quote = Quote.objects.filter(id=row.content_id).first()
        return ApiResponse(data=self._serialize_report(
            row,
            {row.user_id: user} if user else {},
            {row.content_id: quote} if quote else {},
        ))

    @extend_schema(summary=_('AI 安全审计列表'))
    @action(detail=False, methods=['get'], url_path='safety-audits')
    def safety_audits(self, request):
        """AI 安全审计日志列表（输入/输出命中均记一条；需 PERM_OPS）。"""
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        qs = AiSafetyAudit.objects.all().order_by('-id')
        category = (request.query_params.get('category') or '').strip()
        if category:
            qs = qs.filter(safety_category=category)
        app_uids = _user_ids_for_scope_query(request)
        if app_uids is not None:
            qs = qs.filter(user_id__in=app_uids)
        q = (request.query_params.get('q') or '').strip()
        if q:
            user_ids = list(
                User.objects.filter(
                    Q(display_name__icontains=q) | Q(device_id__icontains=q) | Q(username__icontains=q)
                ).values_list('id', flat=True)[:500]
            )
            q_filter = Q(user_id__in=user_ids) | Q(safety_category__icontains=q)
            if q.isdigit():
                q_filter = q_filter | Q(id=int(q)) | Q(user_id=int(q))
            qs = qs.filter(q_filter)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        uids = {r.user_id for r in page}
        user_map = {u.id: u for u in User.objects.filter(id__in=uids)}
        results = []
        for row in page:
            u = user_map.get(row.user_id)
            results.append({
                'id': row.id,
                'user_id': row.user_id,
                'display_name': (u.display_name if u else '') or (u.username if u else '') or '',
                'device_id': u.device_id if u else '',
                'app_key': (getattr(u, 'app_key', '') if u else '') or '',
                'safety_category': row.safety_category,
                'created_at': row.created_at.isoformat() if row.created_at else None,
            })
        return paginator.get_paginated_response(results)

    # ---------- 数据看板统计 ----------

    def _stats_filters(self, request):
        """从请求解析统计筛选条件（日期范围 + language/country/platform/app_key）。"""
        return admin_stats.parse_filters(request.query_params)

    def _require_ops(self, request):
        """校验 PERM_OPS；无权限返回 403 响应，有权限返回 None。"""
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message=_('无权限'), code=403)
        return None

    def _require_content(self, request):
        """校验 PERM_CONTENT。"""
        if not has_perm(request.user, PERM_CONTENT):
            return ApiResponse(message=_('无权限'), code=403)
        return None

    def _require_app_users(self, request):
        """校验 PERM_APP_USERS。"""
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message=_('无权限'), code=403)
        return None

    @extend_schema(summary=_('看板总览 KPI + 序列'))
    @action(detail=False, methods=['get'], url_path='stats/overview')
    def stats_overview(self, request):
        """运营总览指标（所有 stats/* 均需 PERM_OPS）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.overview(d_from, d_to, lang, country, platform, app_key))

    @extend_schema(summary=_('每日 DAU'))
    @action(detail=False, methods=['get'], url_path='stats/dau')
    def stats_dau(self, request):
        """DAU 趋势（基于 last_active_at / 埋点活跃）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.series_dau(d_from, d_to, lang, country, platform, app_key))

    @extend_schema(summary=_('每日新增用户'))
    @action(detail=False, methods=['get'], url_path='stats/new-users')
    def stats_new_users(self, request):
        """新注册用户趋势（按 User.created_at 日聚合）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.series_new_users(d_from, d_to, lang, country, platform, app_key))

    @extend_schema(summary=_('0-30 天留存'))
    @action(detail=False, methods=['get'], url_path='stats/retention')
    def stats_retention(self, request):
        """0-30 天留存表（cohort 相对注册日）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.retention_table(d_from, d_to, lang, country, platform, app_key))

    @extend_schema(summary=_('每日成交总额'))
    @action(detail=False, methods=['get'], url_path='stats/revenue')
    def stats_revenue(self, request):
        """收入趋势（IAP amount_cents，试用通常为 0）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.series_revenue(d_from, d_to, lang, country, platform, app_key))

    @extend_schema(summary=_('地区与设备分布'))
    @action(detail=False, methods=['get'], url_path='stats/geo-device')
    def stats_geo_device(self, request):
        """国家/设备分布（register_country × platform）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        _, _, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.geo_device_dist(lang, country, platform, app_key))

    @extend_schema(summary=_('语言分布'))
    @action(detail=False, methods=['get'], url_path='stats/language')
    def stats_language(self, request):
        """语言分布（system_lang）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        _, _, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.language_dist(lang, country, platform, app_key))

    @extend_schema(summary=_('每日订单/付费率/ARPPU'))
    @action(detail=False, methods=['get'], url_path='stats/orders-kpi')
    def stats_orders_kpi(self, request):
        """每日订单数/付费率/ARPPU。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.series_orders_kpi(d_from, d_to, lang, country, platform, app_key))

    @extend_schema(summary=_('页面/按钮事件 UV·PV'))
    @action(detail=False, methods=['get'], url_path='stats/events')
    def stats_events(self, request):
        """页面/按钮事件 UV·PV 聚合（AnalyticsEvent）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        return ApiResponse(data=admin_stats.events_aggregate(d_from, d_to, lang, country, platform, app_key))

    @extend_schema(summary=_('Manna 功课留存'))
    @action(detail=False, methods=['get'], url_path='stats/manna-practice')
    def stats_manna_practice(self, request):
        """功课 DAU / D1·D7 功课留存 / Quiz 通关与再战（perm=ops）。"""
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        if not app_key:
            app_key = 'manna'
        return ApiResponse(data=admin_stats.manna_practice(d_from, d_to, lang, country, platform, app_key))

    @extend_schema(summary=_('Manna 内容打开次数'))
    @action(detail=False, methods=['get'], url_path='stats/manna-content-opens')
    def stats_manna_content_opens(self, request):
        """按 content_type×content_id 聚合打开 PV/UV（埋点 content_open；perm=ops）。

        Query:
            from / to / app_key / language / country / platform
            content_type: 可选，筛单类型
            limit: Top N，默认 50，最大 200
        """
        denied = self._require_ops(request)
        if denied:
            return denied
        d_from, d_to, lang, country, platform, app_key = self._stats_filters(request)
        if not app_key:
            app_key = 'manna'
        return ApiResponse(data=admin_stats.manna_content_opens(
            d_from, d_to, lang, country, platform, app_key,
            content_type=(request.query_params.get('content_type') or '').strip(),
            limit=request.query_params.get('limit') or 50,
        ))

    # ---------- 用户资产（肯定语 / 收藏）与肯定语模板 ----------

    @extend_schema(summary=_('用户肯定语列表'))
    @action(detail=False, methods=['get'], url_path='user-affirmations')
    def user_affirmations(self, request):
        """分页检索用户自写肯定语（审核用；可按 app_key 过滤用户）。"""
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message='forbidden', code=403)
        qs = Affirmation.objects.all().order_by('-id')
        uid = request.query_params.get('user_id')
        if uid:
            qs = qs.filter(user_id=int(uid))
        app_uids = _user_ids_for_scope_query(request)
        if app_uids is not None:
            qs = qs.filter(user_id__in=app_uids)
        q = (request.query_params.get('q') or '').strip()
        if q:
            qs = qs.filter(Q(content__icontains=q))
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        uids = {r.user_id for r in page}
        users = {
            u.id: u for u in User.objects.filter(id__in=uids).only(
                'id', 'username', 'display_name', 'device_id', 'app_key',
            )
        }
        rows = []
        for r in page:
            u = users.get(r.user_id)
            rows.append({
                'id': r.id,
                'user_id': r.user_id,
                'username': getattr(u, 'username', '') or '',
                'display_name': getattr(u, 'display_name', '') or '',
                'device_id': getattr(u, 'device_id', '') or '',
                'app_key': getattr(u, 'app_key', '') or '',
                'content': r.content,
                'created_at': r.created_at,
                'updated_at': r.updated_at,
            })
        return paginator.get_paginated_response(rows)

    @extend_schema(summary=_('删除用户肯定语'))
    @action(detail=False, methods=['post'], url_path='user-affirmation-delete')
    def user_affirmation_delete(self, request):
        """运营审核删除单条用户肯定语（级联清喜欢/收藏/关注）。"""
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message='forbidden', code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        aid = ser.validated_data['id']
        Affirmation.objects.filter(id=aid).delete()
        UserLike.objects.filter(content_id=aid, target_type='affirmation').delete()
        CollectionItem.objects.filter(content_id=aid, target_type='affirmation').delete()
        AffirmationFollow.objects.filter(affirmation_id=aid).delete()
        return ApiResponse(message='ok')

    @extend_schema(summary=_('用户收藏夹/条目'))
    @action(detail=False, methods=['get'], url_path='user-collections')
    def user_collections(self, request):
        """用户收藏（按 kind × app 区分）。

        kind:
          - folder：语录收藏夹（可再传 folder_id 钻取条目）
          - like：Selah 语录喜欢
          - favorite：Manna 收藏（target_type / target_key）
          - highlight：圣经高亮
          - note：圣经笔记
        app_key：folder/like 按用户 app_key；manna 三类按记录自身 app_key。
        """
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message='forbidden', code=403)
        kind = (request.query_params.get('kind') or 'folder').strip().lower()
        if kind not in ('folder', 'like', 'favorite', 'highlight', 'note'):
            return ApiResponse(message='invalid kind', code=400)
        folder_id = request.query_params.get('folder_id')
        uid = request.query_params.get('user_id')
        q = (request.query_params.get('q') or '').strip()
        app_key = normalize_app_key(request.query_params.get('app_key') or '')
        lang = normalize_lang(
            request.query_params.get('language')
            or request.query_params.get('system_lang')
            or ''
        )
        plat = normalize_platform(request.query_params.get('platform') or '')

        def _users_map(uids):
            return {
                u.id: u for u in User.objects.filter(id__in=uids).only(
                    'id', 'username', 'display_name', 'device_id', 'app_key',
                )
            }

        def _user_ids_for_app():
            if not app_key and not lang and not plat:
                return None
            uq = User.objects.filter(role='user')
            if app_key:
                uq = uq.filter(app_key=app_key)
            if lang:
                uq = uq.filter(system_lang=lang)
            if plat:
                uq = uq.filter(platform=plat)
            return list(uq.values_list('id', flat=True))

        if kind == 'folder' and folder_id:
            qs = CollectionItem.objects.filter(folder_id=int(folder_id)).order_by('-id')
            if uid:
                qs = qs.filter(user_id=int(uid))
            app_uids = _user_ids_for_app()
            if app_uids is not None:
                qs = qs.filter(user_id__in=app_uids)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            cids = [r.content_id for r in page]
            quotes = {
                x.id: x for x in Quote.objects.filter(id__in=cids).only(
                    'id', 'content', 'is_deleted', 'app_key',
                )
            }
            users = _users_map({r.user_id for r in page})
            rows = []
            for r in page:
                qq = quotes.get(r.content_id)
                u = users.get(r.user_id)
                rows.append({
                    'id': r.id,
                    'kind': 'folder_item',
                    'app_key': (getattr(u, 'app_key', '') or '') or (getattr(qq, 'app_key', '') or ''),
                    'user_id': r.user_id,
                    'folder_id': r.folder_id,
                    'content_id': r.content_id,
                    'quote_content': (qq.content if qq else '')[:200],
                    'quote_deleted': bool(qq.is_deleted) if qq else True,
                    'created_at': r.created_at,
                })
            return paginator.get_paginated_response(rows)

        if kind == 'folder':
            from django.db.models import Count
            qs = CollectionFolder.objects.all().order_by('-id')
            if uid:
                qs = qs.filter(user_id=int(uid))
            if q:
                qs = qs.filter(Q(name__icontains=q))
            app_uids = _user_ids_for_app()
            if app_uids is not None:
                qs = qs.filter(user_id__in=app_uids)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            fids = [f.id for f in page]
            counts = {
                row['folder_id']: row['c']
                for row in CollectionItem.objects.filter(folder_id__in=fids)
                .values('folder_id').annotate(c=Count('id'))
            }
            users = _users_map({f.user_id for f in page})
            rows = []
            for f in page:
                u = users.get(f.user_id)
                rows.append({
                    'id': f.id,
                    'kind': 'folder',
                    'folder_type': 'default' if f.is_default else 'custom',
                    'app_key': getattr(u, 'app_key', '') or '',
                    'user_id': f.user_id,
                    'username': getattr(u, 'username', '') or '',
                    'display_name': getattr(u, 'display_name', '') or '',
                    'device_id': getattr(u, 'device_id', '') or '',
                    'name': f.name,
                    'is_default': f.is_default,
                    'item_count': counts.get(f.id, 0),
                    'created_at': f.created_at,
                })
            return paginator.get_paginated_response(rows)

        if kind == 'like':
            qs = UserLike.objects.all().order_by('-id')
            if uid:
                qs = qs.filter(user_id=int(uid))
            app_uids = _user_ids_for_app()
            if app_uids is not None:
                qs = qs.filter(user_id__in=app_uids)
            if q:
                match_ids = list(
                    Quote.objects.filter(content__icontains=q).values_list('id', flat=True)[:2000]
                )
                qs = qs.filter(content_id__in=match_ids)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            cids = [r.content_id for r in page]
            quotes = {
                x.id: x for x in Quote.objects.filter(id__in=cids).only('id', 'content', 'is_deleted')
            }
            users = _users_map({r.user_id for r in page})
            rows = []
            for r in page:
                qq = quotes.get(r.content_id)
                u = users.get(r.user_id)
                rows.append({
                    'id': r.id,
                    'kind': 'like',
                    'app_key': getattr(u, 'app_key', '') or '',
                    'user_id': r.user_id,
                    'username': getattr(u, 'username', '') or '',
                    'display_name': getattr(u, 'display_name', '') or '',
                    'device_id': getattr(u, 'device_id', '') or '',
                    'content_id': r.content_id,
                    'quote_content': (qq.content if qq else '')[:200],
                    'quote_deleted': bool(qq.is_deleted) if qq else True,
                    'created_at': r.created_at,
                })
            return paginator.get_paginated_response(rows)

        model = {'favorite': MannaFavorite, 'highlight': BibleHighlight, 'note': BibleNote}[kind]
        qs = model.objects.all().order_by('-id')
        if uid:
            qs = qs.filter(user_id=int(uid))
        if app_key:
            qs = qs.filter(app_key=app_key)
        # 语言/设备按用户维过滤（记录表无 language 字段）
        if lang or plat:
            uq = User.objects.filter(role='user')
            if lang:
                uq = uq.filter(system_lang=lang)
            if plat:
                uq = uq.filter(platform=plat)
            qs = qs.filter(user_id__in=list(uq.values_list('id', flat=True)))
        if q:
            if kind == 'favorite':
                qs = qs.filter(Q(target_type__icontains=q) | Q(target_key__icontains=q))
            elif kind == 'highlight':
                qs = qs.filter(Q(book_key__icontains=q) | Q(color__icontains=q))
            else:
                qs = qs.filter(Q(book_key__icontains=q) | Q(body__icontains=q))
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        users = _users_map({r.user_id for r in page})
        rows = []
        for r in page:
            u = users.get(r.user_id)
            base = {
                'id': r.id,
                'kind': kind,
                'app_key': getattr(r, 'app_key', '') or getattr(u, 'app_key', '') or '',
                'user_id': r.user_id,
                'username': getattr(u, 'username', '') or '',
                'display_name': getattr(u, 'display_name', '') or '',
                'device_id': getattr(u, 'device_id', '') or '',
                'created_at': r.created_at,
            }
            if kind == 'favorite':
                base.update({
                    'target_type': r.target_type,
                    'target_key': r.target_key,
                    'summary': f'{r.target_type}:{r.target_key}',
                })
            elif kind == 'highlight':
                base.update({
                    'book_key': r.book_key,
                    'chapter': r.chapter,
                    'verse': r.verse,
                    'color': r.color or '',
                    'translation_code': r.translation_code or '',
                    'summary': f'{r.book_key} {r.chapter}:{r.verse}',
                })
            else:
                base.update({
                    'book_key': r.book_key,
                    'chapter': r.chapter,
                    'verse': r.verse,
                    'body': (r.body or '')[:200],
                    'translation_code': r.translation_code or '',
                    'summary': f'{r.book_key} {r.chapter}:{r.verse}',
                })
            rows.append(base)
        return paginator.get_paginated_response(rows)

    @extend_schema(summary=_('删除用户收藏条目'))
    @action(detail=False, methods=['post'], url_path='user-collection-item-delete')
    def user_collection_item_delete(self, request):
        """删除单条：kind=folder_item|like|favorite|highlight|note（默认 folder_item）。"""
        if not has_perm(request.user, PERM_APP_USERS):
            return ApiResponse(message='forbidden', code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        rid = ser.validated_data['id']
        kind = (request.data.get('kind') or 'folder_item').strip().lower()
        model_map = {
            'folder_item': CollectionItem,
            'like': UserLike,
            'favorite': MannaFavorite,
            'highlight': BibleHighlight,
            'note': BibleNote,
        }
        model = model_map.get(kind)
        if not model:
            return ApiResponse(message='invalid kind', code=400)
        model.objects.filter(id=rid).delete()
        return ApiResponse(message='ok')

    @extend_schema(summary=_('肯定语模板'))
    @action(detail=False, methods=['get', 'post'], url_path='affirmation-templates')
    def affirmation_templates(self, request):
        """运营预置肯定语模板列表（GET）或保存/删除（POST）。"""
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message='forbidden', code=403)
        if request.method == 'GET':
            qs = AffirmationTemplate.objects.filter(is_deleted=False).order_by('-sort_weight', '-id')
            qs = _apply_scope_filters(qs, request)
            q = (request.query_params.get('q') or '').strip()
            if q:
                qs = qs.filter(content__icontains=q)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([
                {
                    'id': r.id, 'content': r.content, 'sort_weight': r.sort_weight,
                    'status': r.status,
                    'app_key': r.app_key or '',
                    'language': r.language or '',
                    'platform': r.platform or '',
                    'updated_at': r.updated_at,
                }
                for r in page
            ])
        ser = AffirmationTemplateSaveSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        if data.get('delete_id'):
            row = AffirmationTemplate.objects.filter(id=data['delete_id']).first()
            if row:
                denied = _deny_app_write(request, row.app_key or '')
                if denied:
                    return denied
                AffirmationTemplate.objects.filter(id=row.id).update(is_deleted=True)
            return ApiResponse(message='deleted')
        app = normalize_app_key(data.get('app_key') or '')
        lang = normalize_lang(data.get('language') or '')
        plat = normalize_platform(data.get('platform') or '')
        if data.get('id'):
            existing = AffirmationTemplate.objects.filter(id=data['id'], is_deleted=False).first()
            if existing:
                denied = _deny_app_write(request, existing.app_key or '')
                if denied:
                    return denied
            denied = _deny_app_write(request, app)
            if denied:
                return denied
            patch = {}
            for k in ('content', 'sort_weight', 'status', 'language', 'platform', 'app_key'):
                if k in request.data:
                    if k == 'language':
                        patch[k] = lang
                    elif k == 'platform':
                        patch[k] = plat
                    elif k == 'app_key':
                        patch[k] = app
                    else:
                        patch[k] = data.get(k)
            if patch:
                AffirmationTemplate.objects.filter(id=data['id'], is_deleted=False).update(**patch)
            return ApiResponse(data={'id': data['id']})
        denied = _guard_scoped_write(request, target_app=app, require_app_key=True)
        if denied:
            return denied
        r = AffirmationTemplate.objects.create(
            content=data['content'],
            sort_weight=int(data.get('sort_weight') or 0),
            status=int(data.get('status') if data.get('status') is not None else 1),
            app_key=app,
            language=lang,
            platform=plat,
        )
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('肯定语模板删除'))
    @action(detail=False, methods=['post'], url_path='affirmation-template-delete')
    def affirmation_template_delete(self, request):
        """软删除肯定语模板。"""
        if not has_perm(request.user, PERM_OPS):
            return ApiResponse(message='forbidden', code=403)
        ser = IdSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        row = AffirmationTemplate.objects.filter(id=ser.validated_data['id']).first()
        if row:
            denied = _deny_app_write(request, row.app_key or '')
            if denied:
                return denied
            AffirmationTemplate.objects.filter(id=row.id).update(is_deleted=True)
        return ApiResponse(message='ok')

    # ---------- Manna 内容 / 运营 ----------

    def _manna_scope_norm(self, data):
        return (
            normalize_app_key(data.get('app_key') or ''),
            normalize_lang(data.get('language') or ''),
            normalize_platform(data.get('platform') or ''),
        )

    def _crud_delete_or_patch(
        self, model, request, fields, *, require_fn, extra_create=None, serialize_row=None,
        order_by=('-id',), require_app_key=False, unique_on=None,
    ):
        """通用 GET 分页列表 + POST delete_id / id patch / 新建（含 Staff app_keys）。

        unique_on: 可选唯一维字段元组。patch 若改维撞到已有行，则把本次改动合并到已有行并删当前行
        （典型：语言/设备从具体值改回「全部」时撞上种子全局行）。
        """
        denied = require_fn(request)
        if denied:
            return denied
        if request.method == 'GET':
            qs = model.objects.all().order_by(*order_by)
            if hasattr(model, 'app_key'):
                qs = _apply_scope_filters(qs, request)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            rows = []
            for r in page:
                if serialize_row:
                    rows.append(serialize_row(r))
                else:
                    item = {'id': r.id}
                    for f in fields:
                        val = getattr(r, f, None)
                        if hasattr(val, 'isoformat'):
                            val = val.isoformat()
                        item[f] = val
                    rows.append(item)
            return paginator.get_paginated_response(rows)
        data = request.data if isinstance(request.data, dict) else {}
        if data.get('delete_id'):
            row = model.objects.filter(id=data['delete_id']).first()
            if row and hasattr(row, 'app_key'):
                denied = _deny_app_write(request, getattr(row, 'app_key', '') or '')
                if denied:
                    return denied
            if row:
                row.delete()
            return ApiResponse(message='deleted')
        app, lang, plat = self._manna_scope_norm(data)
        if data.get('id'):
            existing = model.objects.filter(id=data['id']).first()
            if existing and hasattr(existing, 'app_key'):
                denied = _deny_app_write(request, getattr(existing, 'app_key', '') or '')
                if denied:
                    return denied
            if 'app_key' in request.data:
                denied = _deny_app_write(request, app)
                if denied:
                    return denied
            if not existing:
                return ApiResponse(message='not found', code=404)
            patch = {}
            for k in fields:
                if k in request.data:
                    if k == 'language':
                        patch[k] = lang
                    elif k == 'platform':
                        patch[k] = plat
                    elif k == 'app_key':
                        patch[k] = app
                    else:
                        patch[k] = data.get(k)
            if patch:
                merge_target = None
                if unique_on:
                    lookup = {}
                    for f in unique_on:
                        lookup[f] = patch[f] if f in patch else getattr(existing, f, '')
                    merge_target = (
                        model.objects.filter(**lookup).exclude(id=existing.id).first()
                    )
                # 仅「具体语言/设备 → 全部」撞键时合并；反向特化仍返回冲突，避免误删全局种子
                collapsing_scope = False
                if merge_target is not None:
                    if 'language' in patch and patch['language'] == '' and (getattr(existing, 'language', '') or '') != '':
                        collapsing_scope = True
                    if 'platform' in patch and patch['platform'] == '' and (getattr(existing, 'platform', '') or '') != '':
                        collapsing_scope = True
                if merge_target is not None and collapsing_scope:
                    merge_patch = {k: v for k, v in patch.items() if k not in unique_on}
                    with transaction.atomic():
                        if merge_patch:
                            model.objects.filter(id=merge_target.id).update(**merge_patch)
                        model.objects.filter(id=existing.id).delete()
                    return ApiResponse(data={
                        'id': merge_target.id,
                        'merged_into_id': merge_target.id,
                        'deleted_id': existing.id,
                    })
                try:
                    model.objects.filter(id=data['id']).update(**patch)
                except IntegrityError:
                    return ApiResponse(
                        message=_('同 App×语言×设备 下已存在相同唯一键，请换 key 或改维度'),
                        code=400,
                    )
            return ApiResponse(data={'id': data['id']})
        if hasattr(model, 'app_key'):
            denied = _deny_app_write(request, app)
            if denied:
                return denied
            if require_app_key and not app:
                return ApiResponse(message=_('请指定 App（禁止空维新建，避免双品牌串台）'), code=400)
        create_kwargs = {}
        for k in fields:
            if k in data:
                create_kwargs[k] = data.get(k)
            elif k == 'app_key':
                create_kwargs[k] = app
            elif k == 'language':
                create_kwargs[k] = lang
            elif k == 'platform':
                create_kwargs[k] = plat
        if 'app_key' in fields:
            create_kwargs['app_key'] = app
        if 'language' in fields:
            create_kwargs['language'] = lang
        if 'platform' in fields:
            create_kwargs['platform'] = plat
        if extra_create:
            create_kwargs.update(extra_create(data))
        try:
            r = model.objects.create(**create_kwargs)
        except IntegrityError:
            return ApiResponse(
                message=_('同 App×语言×设备 下已存在相同唯一键，请换 key 或改维度'),
                code=400,
            )
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('今日经文 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='daily-verses')
    def daily_verses(self, request):
        fields = (
            'verse_text', 'reference', 'prayer_title', 'prayer_body',
            'bible_book', 'chapter', 'verse_start', 'verse_end',
            'wallpaper_id', 'image_url', 'image_url_night', 'active_date', 'sort_order', 'enabled',
            'app_key', 'language', 'platform',
        )
        return self._crud_delete_or_patch(
            DailyVerse, request, fields, require_fn=self._require_content,
            order_by=('sort_order', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('Verse card 主题 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='versecard-themes')
    def versecard_themes(self, request):
        """底栏主题图；同 app×语言×设备 下 theme_key 唯一，image_url 规范化后不可重复。"""
        from tools.manna_biz import (
            find_versecard_theme_url_conflict,
            normalize_versecard_image_url,
        )

        fields = (
            'theme_key', 'image_url', 'sort_order', 'enabled', 'is_default',
            'app_key', 'language', 'platform',
        )
        if request.method == 'POST':
            data = request.data if isinstance(request.data, dict) else {}
            if not data.get('delete_id'):
                app, lang, plat = self._manna_scope_norm(data)
                if not data.get('id') and not app:
                    return ApiResponse(message=_('请指定 App（禁止空维新建，避免双品牌串台）'), code=400)
                image_url = normalize_versecard_image_url(
                    data.get('image_url') if 'image_url' in data else ''
                )
                theme_key = str(data.get('theme_key') or '').strip()
                if data.get('id'):
                    existing = VerseCardTheme.objects.filter(id=data['id']).first()
                    if existing:
                        if 'image_url' not in data:
                            image_url = normalize_versecard_image_url(existing.image_url)
                        if 'theme_key' not in data:
                            theme_key = existing.theme_key
                        if 'app_key' not in data:
                            app = normalize_app_key(existing.app_key or '')
                        if 'language' not in data:
                            lang = normalize_lang(existing.language or '')
                        if 'platform' not in data:
                            plat = normalize_platform(existing.platform or '')
                if data.get('id') or ('theme_key' in data or 'image_url' in data):
                    if not theme_key and not data.get('id'):
                        return ApiResponse(message='theme_key required', code=400)
                    if not image_url and not data.get('id'):
                        return ApiResponse(message='image_url required', code=400)
                    if image_url:
                        conflict = find_versecard_theme_url_conflict(
                            app, lang, plat, image_url,
                            exclude_id=data.get('id'),
                        )
                        if conflict:
                            return ApiResponse(
                                message=_('同范围内已存在相同背景图（theme_key=%(key)s）') % {
                                    'key': conflict.theme_key,
                                },
                                code=400,
                            )
                if 'image_url' in data and image_url:
                    # 写回规范化 URL，避免编码差异绕过去重
                    mutable = request.data
                    if hasattr(mutable, '_mutable'):
                        mutable._mutable = True
                    if isinstance(mutable, dict):
                        mutable['image_url'] = image_url
                    if theme_key and 'theme_key' in data:
                        mutable['theme_key'] = theme_key
        return self._crud_delete_or_patch(
            VerseCardTheme, request, fields, require_fn=self._require_content,
            order_by=('sort_order', 'id'), require_app_key=True,
            unique_on=('theme_key', 'app_key', 'language', 'platform'),
        )

    @extend_schema(summary=_('圣经译本列表/保存'))
    @action(detail=False, methods=['get', 'post'], url_path='bible-translations')
    def bible_translations(self, request):
        fields = ('code', 'title', 'language', 'sort_order', 'enabled', 'app_key', 'platform')
        return self._crud_delete_or_patch(
            BibleTranslation, request, fields, require_fn=self._require_content,
            order_by=('sort_order', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('圣经章导入'))
    @action(detail=False, methods=['post'], url_path='bible-import-chapter')
    def bible_import_chapter(self, request):
        """JSON 导入一章：translation_code, book_key, book_name, chapter_num, verses[{verse,text}], audio_*。"""
        denied = self._require_content(request)
        if denied:
            return denied
        data = request.data if isinstance(request.data, dict) else {}
        code = str(data.get('translation_code') or data.get('translation') or 'web').strip()
        book_key = str(data.get('book_key') or '').strip()
        book_name = str(data.get('book_name') or book_key).strip()
        try:
            chapter_num = int(data.get('chapter_num') or data.get('chapter') or 0)
        except (TypeError, ValueError):
            chapter_num = 0
        verses = data.get('verses') or []
        if not book_key or chapter_num < 1 or not isinstance(verses, list) or not verses:
            return ApiResponse(message='book_key/chapter/verses required', code=400)
        app = normalize_app_key(data.get('app_key') or 'manna')
        lang = normalize_lang(data.get('language') or 'en')
        plat = normalize_platform(data.get('platform') or '')
        denied = _guard_scoped_write(request, target_app=app, require_app_key=True)
        if denied:
            return denied
        tr, _ = BibleTranslation.objects.get_or_create(
            code=code, app_key=app, language=lang, platform=plat,
            defaults={'title': str(data.get('translation_title') or code.upper()), 'enabled': True},
        )
        book, _ = BibleBook.objects.get_or_create(
            translation_id=tr.id, book_key=book_key,
            defaults={
                'name': book_name,
                'order_index': int(data.get('order_index') or 1),
                'testament': str(data.get('testament') or 'ot'),
                'tag_keys': data.get('tag_keys') or [],
            },
        )
        ch, _ = BibleChapter.objects.update_or_create(
            book_id=book.id, chapter_num=chapter_num,
            defaults={
                'audio_url': str(data.get('audio_url') or ''),
                'audio_duration_sec': int(data.get('audio_duration_sec') or 0),
            },
        )
        created = 0
        for v in verses:
            try:
                vnum = int(v.get('verse') or v.get('verse_num') or 0)
            except (TypeError, ValueError, AttributeError):
                continue
            text = str(v.get('text') or '')
            if vnum < 1 or not text:
                continue
            _, was = BibleVerse.objects.update_or_create(
                chapter_id=ch.id, verse_num=vnum, defaults={'text': text},
            )
            created += 1
        return ApiResponse(data={
            'translation_id': tr.id, 'book_id': book.id, 'chapter_id': ch.id,
            'verse_count': created,
        }, code=201)

    @extend_schema(summary=_('从 content/bible 批量导入'))
    @action(detail=False, methods=['post'], url_path='bible-import-content')
    def bible_import_content(self, request):
        """从 vendored JSON 导入。body: version=WEB|KJV, full=bool, books=[], chapters={}。"""
        denied = self._require_content(request)
        if denied:
            return denied
        from tools.bible_import import import_full_text, import_seed_core, import_version
        data = request.data if isinstance(request.data, dict) else {}
        version = str(data.get('version') or 'WEB').upper()
        app = normalize_app_key(data.get('app_key') or 'manna')
        lang = normalize_lang(data.get('language') or 'en')
        denied = _guard_scoped_write(request, target_app=app, require_app_key=True)
        if denied:
            return denied
        if data.get('seed_core'):
            stats = import_seed_core(app_key=app, language=lang)
            return ApiResponse(data=stats, code=201)
        if data.get('full'):
            stats = import_full_text(version, app_key=app, language=lang)
            return ApiResponse(data={version: stats}, code=201)
        books = data.get('books')
        chapters = data.get('chapters')
        stats = import_version(
            version,
            books=books if isinstance(books, list) else None,
            chapters=chapters if isinstance(chapters, dict) else None,
            full=False,
            app_key=app,
            language=lang,
        )
        return ApiResponse(data={version: stats}, code=201)

    @extend_schema(summary=_('书卷列表/改名'))
    @action(detail=False, methods=['get', 'post'], url_path='bible-books')
    def bible_books(self, request):
        denied = self._require_content(request)
        if denied:
            return denied
        if request.method == 'GET':
            try:
                tr_id = int(request.query_params.get('translation_id') or 0)
            except (TypeError, ValueError):
                tr_id = 0
            qs = BibleBook.objects.all().order_by('order_index', 'id')
            if tr_id:
                qs = qs.filter(translation_id=tr_id)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([{
                'id': b.id,
                'translation_id': b.translation_id,
                'book_key': b.book_key,
                'name': b.name,
                'order_index': b.order_index,
                'testament': b.testament,
                'tag_keys': b.tag_keys or [],
            } for b in page])
        data = request.data if isinstance(request.data, dict) else {}
        if not data.get('id'):
            return ApiResponse(message='id required', code=400)
        row = BibleBook.objects.filter(id=data['id']).first()
        if not row:
            return ApiResponse(message='not found', code=404)
        patch = {}
        for k in ('name', 'order_index', 'testament', 'tag_keys'):
            if k in data:
                patch[k] = data.get(k)
        if patch:
            BibleBook.objects.filter(id=row.id).update(**patch)
        return ApiResponse(data={'id': row.id})

    @extend_schema(summary=_('章节列表/改音频'))
    @action(detail=False, methods=['get', 'post'], url_path='bible-chapters')
    def bible_chapters(self, request):
        denied = self._require_content(request)
        if denied:
            return denied
        if request.method == 'GET':
            try:
                book_id = int(request.query_params.get('book_id') or 0)
            except (TypeError, ValueError):
                book_id = 0
            qs = BibleChapter.objects.all().order_by('chapter_num', 'id')
            if book_id:
                qs = qs.filter(book_id=book_id)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([{
                'id': c.id,
                'book_id': c.book_id,
                'chapter_num': c.chapter_num,
                'audio_url': c.audio_url or '',
                'audio_duration_sec': c.audio_duration_sec,
            } for c in page])
        data = request.data if isinstance(request.data, dict) else {}
        if not data.get('id'):
            return ApiResponse(message='id required', code=400)
        row = BibleChapter.objects.filter(id=data['id']).first()
        if not row:
            return ApiResponse(message='not found', code=404)
        patch = {}
        for k in ('audio_url', 'audio_duration_sec'):
            if k in data:
                patch[k] = data.get(k)
        if patch:
            BibleChapter.objects.filter(id=row.id).update(**patch)
        return ApiResponse(data={'id': row.id})

    @extend_schema(summary=_('经文列表/改错'))
    @action(detail=False, methods=['get', 'post'], url_path='bible-verses')
    def bible_verses(self, request):
        denied = self._require_content(request)
        if denied:
            return denied
        if request.method == 'GET':
            try:
                chapter_id = int(request.query_params.get('chapter_id') or 0)
            except (TypeError, ValueError):
                chapter_id = 0
            qs = BibleVerse.objects.all().order_by('verse_num', 'id')
            if chapter_id:
                qs = qs.filter(chapter_id=chapter_id)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([{
                'id': v.id,
                'chapter_id': v.chapter_id,
                'verse_num': v.verse_num,
                'text': v.text,
            } for v in page])
        data = request.data if isinstance(request.data, dict) else {}
        if not data.get('id'):
            return ApiResponse(message='id required', code=400)
        row = BibleVerse.objects.filter(id=data['id']).first()
        if not row:
            return ApiResponse(message='not found', code=404)
        if 'text' in data:
            BibleVerse.objects.filter(id=row.id).update(text=str(data.get('text') or ''))
        return ApiResponse(data={'id': row.id})

    @extend_schema(summary=_('灵修系列 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='devotion-series')
    def devotion_series(self, request):
        fields = (
            'title', 'cover_url', 'pastor_bio', 'pastor_name', 'pastor_avatar_url',
            'tags', 'session_count', 'is_vip',
            'sort_order', 'enabled', 'app_key', 'language', 'platform',
        )
        if request.method == 'POST':
            data = request.data
            if hasattr(data, 'get') and 'pastor_bio' not in data and 'subtitle' in data:
                if hasattr(data, '_mutable'):
                    data._mutable = True
                data['pastor_bio'] = data.get('subtitle')

        return self._crud_delete_or_patch(
            DevotionSeries, request, fields, require_fn=self._require_content,
            order_by=('sort_order', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('灵修课时 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='devotion-sessions')
    def devotion_sessions(self, request):
        denied = self._require_content(request)
        if denied:
            return denied
        if request.method == 'GET':
            qs = DevotionSession.objects.all().order_by('sort_order', 'id')
            sid = request.query_params.get('series_id')
            if sid:
                qs = qs.filter(series_id=sid)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([{
                'id': r.id, 'series_id': r.series_id, 'title': r.title,
                'scripture_ref': r.scripture_ref, 'scripture_text': r.scripture_text,
                'body': r.body,
                'sort_order': r.sort_order, 'duration_hint_sec': r.duration_hint_sec,
                'enabled': r.enabled,
            } for r in page])
        data = request.data if isinstance(request.data, dict) else {}
        if data.get('delete_id'):
            row = DevotionSession.objects.filter(id=data['delete_id']).first()
            if row:
                parent = DevotionSeries.objects.filter(id=row.series_id).first()
                if parent:
                    denied = _deny_app_write(request, parent.app_key or '')
                    if denied:
                        return denied
                series_id = row.series_id
                row.delete()
                _refresh_devotion_session_count(series_id)
            return ApiResponse(message='deleted')
        if data.get('id'):
            existing = DevotionSession.objects.filter(id=data['id']).first()
            if existing:
                parent = DevotionSeries.objects.filter(id=existing.series_id).first()
                if parent:
                    denied = _deny_app_write(request, parent.app_key or '')
                    if denied:
                        return denied
            patch = {k: data[k] for k in (
                'title', 'scripture_ref', 'scripture_text', 'body', 'sort_order',
                'duration_hint_sec', 'enabled', 'series_id',
            ) if k in request.data}
            if patch:
                old_series_id = existing.series_id if existing else None
                DevotionSession.objects.filter(id=data['id']).update(**patch)
                if 'enabled' in patch or 'series_id' in patch:
                    new_series_id = patch.get('series_id', old_series_id)
                    _refresh_devotion_session_count(old_series_id)
                    if new_series_id != old_series_id:
                        _refresh_devotion_session_count(new_series_id)
            return ApiResponse(data={'id': data['id']})
        if not data.get('series_id') or not data.get('title') or not data.get('body'):
            return ApiResponse(message='series_id/title/body required', code=400)
        parent = DevotionSeries.objects.filter(id=int(data['series_id'])).first()
        if not parent:
            return ApiResponse(message='series not found', code=404)
        denied = _deny_app_write(request, parent.app_key or '')
        if denied:
            return denied
        r = DevotionSession.objects.create(
            series_id=int(data['series_id']),
            title=data['title'],
            scripture_ref=data.get('scripture_ref') or '',
            scripture_text=data.get('scripture_text') or '',
            body=data['body'],
            sort_order=int(data.get('sort_order') or 0),
            duration_hint_sec=int(data.get('duration_hint_sec') or 0),
            enabled=bool(data.get('enabled', True)),
        )
        _refresh_devotion_session_count(r.series_id)
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('Manna 测验关卡 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='manna-quiz-levels')
    def manna_quiz_levels(self, request):
        fields = (
            'level_no', 'title', 'time_limit_sec', 'pass_score', 'reward_manna',
            'fail_penalty', 'sort_order', 'enabled', 'app_key', 'language', 'platform',
        )
        return self._crud_delete_or_patch(
            MannaQuizLevel, request, fields, require_fn=self._require_content,
            order_by=('sort_order', 'level_no', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('Manna 测验题目 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='manna-quiz-questions')
    def manna_quiz_questions(self, request):
        denied = self._require_content(request)
        if denied:
            return denied
        if request.method == 'GET':
            qs = MannaQuizQuestion.objects.all().order_by('sort_order', 'id')
            lid = request.query_params.get('level_id')
            if lid:
                qs = qs.filter(level_id=lid)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([{
                'id': r.id, 'level_id': r.level_id, 'qtype': r.qtype, 'stem': r.stem,
                'options': r.options, 'answer_key': r.answer_key, 'bible_ref': r.bible_ref,
                'explanation': r.explanation, 'sort_order': r.sort_order, 'enabled': r.enabled,
            } for r in page])
        data = request.data if isinstance(request.data, dict) else {}
        if data.get('delete_id'):
            row = MannaQuizQuestion.objects.filter(id=data['delete_id']).first()
            if row:
                parent = MannaQuizLevel.objects.filter(id=row.level_id).first()
                if parent:
                    denied = _deny_app_write(request, parent.app_key or '')
                    if denied:
                        return denied
                row.delete()
            return ApiResponse(message='deleted')
        if data.get('id'):
            existing = MannaQuizQuestion.objects.filter(id=data['id']).first()
            if existing:
                parent = MannaQuizLevel.objects.filter(id=existing.level_id).first()
                if parent:
                    denied = _deny_app_write(request, parent.app_key or '')
                    if denied:
                        return denied
            patch = {k: data[k] for k in (
                'level_id', 'qtype', 'stem', 'options', 'answer_key', 'bible_ref',
                'explanation', 'sort_order', 'enabled',
            ) if k in request.data}
            if 'bible_ref' in patch:
                ref, err = _normalize_quiz_bible_ref(patch.get('bible_ref'))
                if err:
                    return ApiResponse(message=err, code=400)
                patch['bible_ref'] = ref
            if patch:
                MannaQuizQuestion.objects.filter(id=data['id']).update(**patch)
            return ApiResponse(data={'id': data['id']})
        if not data.get('level_id') or not data.get('stem') or not data.get('answer_key'):
            return ApiResponse(message='level_id/stem/answer_key required', code=400)
        parent = MannaQuizLevel.objects.filter(id=int(data['level_id'])).first()
        if not parent:
            return ApiResponse(message='level not found', code=404)
        denied = _deny_app_write(request, parent.app_key or '')
        if denied:
            return denied
        ref, err = _normalize_quiz_bible_ref(data.get('bible_ref'))
        if err:
            return ApiResponse(message=err, code=400)
        r = MannaQuizQuestion.objects.create(
            level_id=int(data['level_id']),
            qtype=str(data.get('qtype') or 'single'),
            stem=data['stem'],
            options=data.get('options') or [],
            answer_key=str(data['answer_key']),
            bible_ref=ref,
            explanation=str(data.get('explanation') or ''),
            sort_order=int(data.get('sort_order') or 0),
            enabled=bool(data.get('enabled', True)),
        )
        return ApiResponse(data={'id': r.id}, code=201)

    @extend_schema(summary=_('Manna 故事 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='manna-stories')
    def manna_stories(self, request):
        denied = self._require_content(request)
        if denied:
            return denied
        if request.method == 'GET':
            qs = MannaQuizStory.objects.all().order_by('sort_order', 'id')
            qs = _apply_scope_filters(qs, request)
            paginator = CustomPagination()
            page = paginator.paginate_queryset(qs, request)
            return paginator.get_paginated_response([_serialize_manna_story(r) for r in page])

        data = request.data if isinstance(request.data, dict) else {}
        if data.get('delete_id'):
            row = MannaQuizStory.objects.filter(id=data['delete_id']).first()
            if row:
                denied = _deny_app_write(request, getattr(row, 'app_key', '') or '')
                if denied:
                    return denied
                row.delete()
            return ApiResponse(message='deleted')

        app, lang, plat = self._manna_scope_norm(data)
        simple_fields = (
            'title', 'unlock_cost', 'knowledge_cards', 'linked_level_id',
            'sort_order', 'enabled',
        )

        if data.get('id'):
            existing = MannaQuizStory.objects.filter(id=data['id']).first()
            if not existing:
                return ApiResponse(message='not found', code=404)
            denied = _deny_app_write(request, getattr(existing, 'app_key', '') or '')
            if denied:
                return denied
            if 'app_key' in request.data:
                denied = _deny_app_write(request, app)
                if denied:
                    return denied
            patch = {}
            for k in simple_fields:
                if k in request.data:
                    patch[k] = data.get(k)
            if 'language' in request.data:
                patch['language'] = lang
            if 'platform' in request.data:
                patch['platform'] = plat
            if 'app_key' in request.data:
                patch['app_key'] = app
            # 未传 blocks 时仍允许手改 cover/card/body（兼容旧运营）；传了则派生覆盖
            if 'blocks' not in request.data:
                for k in ('cover_url', 'card_url', 'body'):
                    if k in request.data:
                        patch[k] = data.get(k)
            err = _apply_story_blocks_fields(
                patch, request.data, existing=existing, require_image=False,
            )
            if err:
                return ApiResponse(message=err, code=400)
            if patch:
                MannaQuizStory.objects.filter(id=data['id']).update(**patch)
            return ApiResponse(data={'id': data['id']})

        denied = _deny_app_write(request, app)
        if denied:
            return denied
        if not app:
            return ApiResponse(message=_('请指定 App（禁止空维新建，避免双品牌串台）'), code=400)
        if not (data.get('title') or '').strip():
            return ApiResponse(message='title required', code=400)

        create_kwargs = {
            'title': str(data.get('title') or '').strip(),
            'unlock_cost': int(data.get('unlock_cost') or 0),
            'knowledge_cards': data.get('knowledge_cards') or [],
            'linked_level_id': data.get('linked_level_id'),
            'sort_order': int(data.get('sort_order') or 0),
            'enabled': bool(data.get('enabled', True)),
            'app_key': app,
            'language': lang,
            'platform': plat,
            'cover_url': str(data.get('cover_url') or ''),
            'card_url': str(data.get('card_url') or ''),
            'body': str(data.get('body') or ''),
            'blocks': [],
            'cover_image_index': 0,
        }
        try:
            if create_kwargs['linked_level_id'] in ('', None):
                create_kwargs['linked_level_id'] = None
            elif create_kwargs['linked_level_id'] is not None:
                create_kwargs['linked_level_id'] = int(create_kwargs['linked_level_id'])
        except (TypeError, ValueError):
            return ApiResponse(message='linked_level_id invalid', code=400)

        patch = {}
        err = _apply_story_blocks_fields(
            patch, data, existing=None,
            require_image=bool(data.get('blocks')),
        )
        if err:
            return ApiResponse(message=err, code=400)
        create_kwargs.update(patch)
        row = MannaQuizStory.objects.create(**create_kwargs)
        return ApiResponse(data={'id': row.id}, code=201)

    @extend_schema(summary=_('祷告内容 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='prayer-contents')
    def prayer_contents(self, request):
        fields = (
            'kind', 'title', 'body', 'verse_text', 'reference', 'image_url', 'version',
            'sort_order', 'enabled', 'app_key', 'language', 'platform',
        )
        return self._crud_delete_or_patch(
            PrayerContent, request, fields, require_fn=self._require_ops,
            order_by=('kind', 'sort_order', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('徽章定义 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='badge-defs')
    def badge_defs(self, request):
        fields = (
            'code', 'title', 'icon_url', 'rule', 'sort_order', 'enabled',
            'app_key', 'language', 'platform',
        )
        return self._crud_delete_or_patch(
            BadgeDef, request, fields, require_fn=self._require_ops,
            order_by=('sort_order', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('Home 运营卡 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='manna-home-cards')
    def manna_home_cards(self, request):
        fields = (
            'card_type', 'title', 'subtitle', 'cover_url', 'link_path', 'payload',
            'sort_order', 'enabled', 'active_date', 'app_key', 'language', 'platform',
        )
        return self._crud_delete_or_patch(
            MannaHomeCard, request, fields, require_fn=self._require_content,
            order_by=('sort_order', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('Testimony CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='manna-testimonies')
    def manna_testimonies(self, request):
        fields = (
            'title', 'body', 'author_label', 'cover_url', 'sort_order', 'enabled',
            'app_key', 'language', 'platform',
        )
        return self._crud_delete_or_patch(
            MannaTestimony, request, fields, require_fn=self._require_content,
            order_by=('sort_order', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('成长卡文案 CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='manna-stats-copies')
    def manna_stats_copies(self, request):
        fields = ('line', 'language', 'app_key', 'platform', 'enabled', 'sort_order')
        return self._crud_delete_or_patch(
            MannaStatsCopy, request, fields, require_fn=self._require_ops,
            order_by=('sort_order', 'id'), require_app_key=True,
        )

    @extend_schema(summary=_('内容 Tag CRUD'))
    @action(detail=False, methods=['get', 'post'], url_path='manna-tags')
    def manna_tags(self, request):
        fields = (
            'key', 'display_name', 'featured_series_id', 'sort_order', 'enabled',
            'app_key', 'language', 'platform',
        )
        return self._crud_delete_or_patch(
            MannaContentTag, request, fields, require_fn=self._require_content,
            order_by=('sort_order', 'id'), require_app_key=True,
            unique_on=('key', 'app_key', 'language', 'platform'),
        )

    @extend_schema(summary=_('按 Tag 配置灵修推荐书目'))
    @action(detail=False, methods=['get', 'post'], url_path='manna-tag-featured')
    def manna_tag_featured(self, request):
        """运营配置某 Tag 下的推荐灵修系列。

        GET  ?app_key=&language=&platform=&tag=  列出 Tag 及当前推荐
        POST {tag_id|tag_key, featured_series_id|null, app_key, language?, platform?}
             featured_series_id 传 null/0/空 表示清除运营推荐（回退自动挑选）
        """
        denied = self._require_content(request)
        if denied:
            return denied

        if request.method == 'GET':
            app = normalize_app_key(request.query_params.get('app_key') or 'manna') or 'manna'
            lang = normalize_lang(request.query_params.get('language') or '')
            plat = normalize_platform(request.query_params.get('platform') or '')
            tag_q = (request.query_params.get('tag') or '').strip()
            qs = MannaContentTag.objects.filter(enabled=True, app_key=app).order_by('sort_order', 'id')
            if lang:
                qs = qs.filter(Q(language='') | Q(language=lang))
            if plat:
                qs = qs.filter(Q(platform='') | Q(platform=plat))
            if tag_q:
                qs = qs.filter(Q(key__iexact=tag_q) | Q(display_name__iexact=tag_q))
            rows = list(qs)
            series_ids = [int(r.featured_series_id) for r in rows if r.featured_series_id]
            series_map = {
                s.id: s for s in DevotionSeries.objects.filter(id__in=series_ids)
            } if series_ids else {}
            data = []
            for r in rows:
                sid = r.featured_series_id
                try:
                    sid = int(sid) if sid is not None else None
                except (TypeError, ValueError):
                    sid = None
                series = series_map.get(sid) if sid else None
                data.append({
                    'tag_id': r.id,
                    'key': r.key,
                    'display_name': r.display_name,
                    'featured_series_id': sid,
                    'featured_series_title': series.title if series else '',
                    'featured_series_enabled': bool(series.enabled) if series else False,
                    'sort_order': r.sort_order,
                    'app_key': r.app_key,
                    'language': r.language,
                    'platform': r.platform,
                })
            return ApiResponse(data=data)

        data = request.data if isinstance(request.data, dict) else {}
        app, lang, plat = self._manna_scope_norm(data)
        denied = _deny_app_write(request, app)
        if denied:
            return denied
        if not app:
            return ApiResponse(message=_('请指定 App（禁止空维新建，避免双品牌串台）'), code=400)

        tag_row = None
        tag_id = data.get('tag_id') or data.get('id')
        if tag_id:
            try:
                tag_row = MannaContentTag.objects.filter(id=int(tag_id)).first()
            except (TypeError, ValueError):
                tag_row = None
        if not tag_row:
            tag_key = (data.get('tag_key') or data.get('tag') or data.get('key') or '').strip()
            if not tag_key:
                return ApiResponse(message='tag_id or tag_key required', code=400)
            tag_row = (
                MannaContentTag.objects.filter(app_key=app)
                .filter(Q(key__iexact=tag_key) | Q(display_name__iexact=tag_key))
                .filter(Q(language=lang) | Q(language=''))
                .filter(Q(platform=plat) | Q(platform=''))
                .order_by('-language', '-platform', 'id')
                .first()
            )
        if not tag_row:
            return ApiResponse(message='tag not found', code=404)
        denied = _deny_app_write(request, getattr(tag_row, 'app_key', '') or '')
        if denied:
            return denied

        raw_fid = data.get('featured_series_id', ...)
        if raw_fid is ...:
            return ApiResponse(message='featured_series_id required (use null to clear)', code=400)
        if raw_fid in (None, '', 0, '0', 'null'):
            featured_series_id = None
        else:
            try:
                featured_series_id = int(raw_fid)
            except (TypeError, ValueError):
                return ApiResponse(message='invalid featured_series_id', code=400)
            series = DevotionSeries.objects.filter(id=featured_series_id, enabled=True).first()
            if not series:
                return ApiResponse(message='series not found or disabled', code=404)

        tag_row.featured_series_id = featured_series_id
        tag_row.save(update_fields=['featured_series_id', 'updated_at'])
        series = (
            DevotionSeries.objects.filter(id=featured_series_id).first()
            if featured_series_id else None
        )
        return ApiResponse(data={
            'tag_id': tag_row.id,
            'key': tag_row.key,
            'display_name': tag_row.display_name,
            'featured_series_id': featured_series_id,
            'featured_series_title': series.title if series else '',
        })

    @extend_schema(summary=_('Manna 流水只读'))
    @action(detail=False, methods=['get'], url_path='manna-ledger')
    def manna_ledger_admin(self, request):
        denied = self._require_app_users(request)
        if denied:
            return denied
        try:
            uid = int(request.query_params.get('user_id') or 0)
        except (TypeError, ValueError):
            uid = 0
        if not uid:
            # 无 user_id：按日 × reason 汇总（功课看板同源）
            app = normalize_app_key(request.query_params.get('app_key') or 'manna') or 'manna'
            reason = str(request.query_params.get('reason') or '').strip()
            d_from, d_to, _, _, _, _ = self._stats_filters(request)
            start = datetime.combine(d_from, datetime.min.time(), tzinfo=dt_tz.utc)
            end = datetime.combine(d_to + timedelta(days=1), datetime.min.time(), tzinfo=dt_tz.utc)
            qs = MannaLedger.objects.filter(app_key=app, created_at__gte=start, created_at__lt=end)
            if reason:
                qs = qs.filter(reason=reason)
            from django.db.models.functions import TruncDate
            from django.db.models import Sum, Count
            rows = list(
                qs.annotate(day=TruncDate('created_at'))
                .values('day', 'reason')
                .annotate(delta_sum=Sum('delta'), count=Count('id'))
                .order_by('-day', 'reason')
            )
            return ApiResponse(data=[{
                'date': r['day'].isoformat() if r['day'] else None,
                'reason': r['reason'],
                'delta_sum': int(r['delta_sum'] or 0),
                'count': int(r['count'] or 0),
            } for r in rows])
        app = normalize_app_key(request.query_params.get('app_key') or 'manna')
        qs = MannaLedger.objects.filter(user_id=uid, app_key=app).order_by('-created_at', '-id')
        reason = str(request.query_params.get('reason') or '').strip()
        if reason:
            qs = qs.filter(reason=reason)
        paginator = CustomPagination()
        page = paginator.paginate_queryset(qs, request)
        return paginator.get_paginated_response([{
            'id': r.id, 'delta': r.delta, 'balance_after': r.balance_after,
            'reason': r.reason, 'ref_type': r.ref_type, 'ref_id': r.ref_id,
            'remark': r.remark,
            'created_at': r.created_at.isoformat() if r.created_at else None,
        } for r in page])

